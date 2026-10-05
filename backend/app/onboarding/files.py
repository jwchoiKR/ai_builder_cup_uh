"""Onboarding step 1: store uploaded files immutably, and turn each into what the model can read whole.

PDF, images, plain text and HTML go to Gemini as they are. DOCX and XLSX/CSV, which Gemini does not read natively,
become text in code: DOCX paragraphs and table rows, spreadsheet rows with their sheet name and row number (so
evidence can point at "sheet Prices row 7").
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import mimetypes
from dataclasses import dataclass
from pathlib import Path

NATIVE = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
          ".webp": "image/webp", ".txt": "text/plain", ".md": "text/plain", ".html": "text/html",
          ".htm": "text/html"}
CONVERTED = {".docx", ".xlsx", ".csv"}
MAX_BYTES = 20 * 1024 * 1024  # per file


@dataclass(frozen=True)
class StoredFile:
    name: str
    sha256: str
    path: Path
    size: int


@dataclass(frozen=True)
class ModelPart:
    name: str
    data: bytes
    mime: str


class UnsupportedFile(ValueError):
    pass


def store(root: Path, draft_id: str, name: str, data: bytes) -> StoredFile:
    """Content-addressed and never overwritten: data/onboarding/<draft_id>/<sha256><ext>, plus an index of names."""
    if len(data) > MAX_BYTES:
        raise UnsupportedFile(f"{name}: larger than {MAX_BYTES // (1024 * 1024)} MB")
    ext = Path(name).suffix.lower()
    if ext not in NATIVE and ext not in CONVERTED:
        raise UnsupportedFile(f"{name}: file type {ext or '(none)'} is not supported")
    sha = hashlib.sha256(data).hexdigest()
    folder = root / draft_id
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{sha}{ext}"
    if not path.exists():
        path.write_bytes(data)
    index = folder / "files.json"
    entries = json.loads(index.read_text(encoding="utf-8")) if index.exists() else []
    if not any(e["sha256"] == sha and e["name"] == name for e in entries):
        entries.append({"name": name, "sha256": sha, "file": path.name, "bytes": len(data)})
        index.write_text(json.dumps(entries, ensure_ascii=False, indent=1), encoding="utf-8")
    return StoredFile(name, sha, path, len(data))


def _docx_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    out = [f"[paragraph {i + 1}] {p.text}" for i, p in enumerate(d.paragraphs) if p.text.strip()]
    for t, table in enumerate(d.tables, 1):
        for r, row in enumerate(table.rows, 1):
            out.append(f"[table {t} row {r}] " + " | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def _xlsx_text(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                out.append(f"[sheet {ws.title} row {r}] " + " | ".join(cells))
    return "\n".join(out)


def _csv_text(data: bytes) -> str:
    rows = csv.reader(io.StringIO(data.decode("utf-8-sig")))
    return "\n".join(f"[row {r}] " + " | ".join(row) for r, row in enumerate(rows, 1) if any(row))


def model_part(name: str, data: bytes) -> ModelPart:
    ext = Path(name).suffix.lower()
    if ext in NATIVE:
        return ModelPart(name, data, NATIVE[ext])
    text = {".docx": _docx_text, ".xlsx": _xlsx_text, ".csv": _csv_text}[ext](data)
    return ModelPart(name, f"FILE {name} (converted to text)\n{text}".encode(), "text/plain")


def media_type(name: str) -> str:
    return NATIVE.get(Path(name).suffix.lower()) or mimetypes.guess_type(name)[0] or "application/octet-stream"

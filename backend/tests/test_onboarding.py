"""P0 onboarding steps 1-3: immutable storage, file conversion, the read call and code normalisation."""
from __future__ import annotations

import io
import json

import pytest
from pathlib import Path

from app.onboarding.files import UnsupportedFile, model_part, store
from app.onboarding.read import normalise, read_product
from app.onboarding.record import ProductRecord

from .helpers import make_ctx
from .judge_fakes import ScriptedBackend

EV = [{"file": "spec.docx", "locator": "paragraph 1", "quote": "P 30 ml, S$36.00; Zinc 0.5 mg"}]


def test_store_is_content_addressed_and_never_overwrites(tmp_path):
    a = store(tmp_path, "d1", "spec.txt", b"one")
    again = store(tmp_path, "d1", "spec.txt", b"one")
    b = store(tmp_path, "d1", "spec.txt", b"two")  # a new version is a new file
    assert a.path == again.path and a.path != b.path and a.path.read_bytes() == b"one"
    with pytest.raises(UnsupportedFile):
        store(tmp_path, "d1", "virus.exe", b"x")


def test_docx_and_xlsx_become_located_text():
    import docx
    import openpyxl

    d = docx.Document()
    d.add_paragraph("Warnings: patch test first.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Niacin", "2%"
    buf = io.BytesIO()
    d.save(buf)
    text = model_part("spec.docx", buf.getvalue()).data.decode()
    assert "[paragraph 1] Warnings: patch test first." in text and "[table 1 row 1] Niacin | 2%" in text
    wb = openpyxl.Workbook()
    wb.active.title = "Prices"
    wb.active.append(["SKU", "Price"])
    wb.active.append(["X-1", "S$10.00"])
    buf = io.BytesIO()
    wb.save(buf)
    part = model_part("prices.xlsx", buf.getvalue())
    assert part.mime == "text/plain" and "[sheet Prices row 2] X-1 | S$10.00" in part.data.decode()
    assert model_part("label.jpg", b"\xff\xd8").mime == "image/jpeg"


def _record(**kw) -> dict:
    base = {"brand": "B", "product_name": "P", "short_names": ["p"], "generic_word": "Serum", "category": "Cosmetic",
            "category_evidence": EV, "language": "en",
            "variants": [{"variant_id": "A", "name": "P 30", "size_value": "30.0 ml", "size_unit": "mL",
                          "market": "sg", "prices": [{"currency": "S$", "minor": 3600, "kind": "live", "market": "sg",
                                                       "evidence": EV},
                                                      {"currency": "??", "minor": 1, "kind": "list", "market": None,
                                                       "evidence": EV}], "evidence": EV},
                         {"variant_id": "A", "name": "P 50", "size_value": None, "size_unit": None, "market": None,
                          "prices": [], "evidence": EV},
                         {"variant_id": "Z", "name": "made up", "size_value": None, "size_unit": None, "market": None,
                          "prices": [], "evidence": []}],
            "offers": [], "ingredients": [{"name": "Zinc", "amount_value": "0.50", "amount_unit": "MG", "purpose": None,
                                           "evidence": EV}],
            "ingredient_list_complete": False, "approved_claims": [], "not_approved_claims": [],
            "avoid_saying": [], "qualified_results": [], "cautions": [{"text": "Patch test.", "evidence": []}],
            "how_to_use": [], "key_facts": [], "missing": [], "conflicts": []}
    return base | kw


def test_normalise_currency_units_identity_and_evidence():
    r, notes = normalise(ProductRecord.model_validate(_record()))
    v = r.variants
    assert [x.variant_id for x in v] == ["A", "A_2"]  # the one without evidence is dropped, ids are unique
    assert (v[0].size_value, v[0].size_unit, v[0].market) == ("30", "ml", "SG")
    assert [(m.currency, m.minor, m.market) for m in v[0].prices] == [("SGD", 3600, "SG")]
    assert (r.ingredients[0].amount_value, r.ingredients[0].amount_unit) == ("0.5", "mg")
    assert r.cautions == [] and r.generic_word == "serum" and r.category == "cosmetic"
    assert any("unknown currency" in n for n in notes) and any("cautions" in n for n in notes)


def test_every_field_is_required_in_the_model_schema():
    schema = ProductRecord.model_json_schema()
    assert set(schema["required"]) == set(schema["properties"])
    assert set(schema["$defs"]["Variant"]["required"]) == set(schema["$defs"]["Variant"]["properties"])


async def test_read_product_sends_every_file_and_normalises(tmp_path):
    backend = ScriptedBackend()
    backend.outputs[("onboarding_read", "spec.txt|label.jpg")] = _record()
    seen = []
    original = backend.generate

    async def capture(req, timeout_ms):
        seen.append(req)
        return await original(req, timeout_ms)

    backend.generate = capture
    ctx = make_ctx(backend=backend)
    try:
        res = await read_product(ctx.llm, [model_part("spec.txt", b"hello"), model_part("label.jpg", b"\xff\xd8")])
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    [req] = [r for r in seen if r.agent == "onboarding_read"]
    assert [m for _, m in req.parts] == ["text/plain", "text/plain", "text/plain", "image/jpeg"]
    assert [r.agent for r in seen if r.agent != "onboarding_read"] == ["onboarding_cautions",  # D-226
                                                                      "onboarding_audit"]  # D-249
    [audit] = [r for r in seen if r.agent == "onboarding_audit"]
    assert audit.parts == req.parts and '<data name="RECORD">' in audit.user_message
    assert res.method == "gemini_whole" and res.record.variants[0].prices[0].currency == "SGD"


def test_d217_the_upload_api_runs_the_same_pipeline(tmp_path):
    import time

    from fastapi.testclient import TestClient

    from app.main import create_app

    from .helpers import make_settings

    backend = ScriptedBackend()
    backend.outputs[("onboarding_read", "spec.txt")] = _record()
    with TestClient(create_app(make_settings(ONBOARDING_DIR=str(tmp_path)), llm_backend=backend)) as c:
        def wait(job: str) -> dict:
            for _ in range(100):
                got = c.get(f"/api/onboarding/jobs/{job}").json()
                if got["status"] in ("done", "failed"):
                    return got
                time.sleep(0.02)
            return got

        r = c.post("/api/onboarding/jobs", files=[("files", ("spec.txt", b"hello", "text/plain"))])
        got = wait(r.json()["job_id"])
        assert got["status"] == "done" and got["result"]["record"]["variants"][0]["prices"][0]["currency"] == "SGD"
        assert (tmp_path / r.json()["job_id"] / "result.json").exists()
        bad = c.post("/api/onboarding/jobs", files=[("files", ("x.exe", b"MZ", "application/octet-stream"))])
        assert wait(bad.json()["job_id"])["status"] == "failed"


async def test_d218_a_release_runs_its_own_one_segment_show(tmp_path):
    from app.onboarding.release import product_key, save_release
    from app.pipelines.runtime import create_session

    rec, _ = normalise(ProductRecord.model_validate(_record(
        cautions=[{"text": "Patch test first.", "evidence": EV}],
        approved_claims=[{"text": "Helps skin feel soft.", "evidence": EV}])))
    rid = save_release(tmp_path, rec, "job1")
    assert save_release(tmp_path, rec, "job2") == rid.replace("_v1", "_v2")  # every approval is a new version
    ctx = make_ctx(ONBOARDING_DIR=str(tmp_path))
    try:
        s = await create_session(ctx, "ignored", release_id=rid)
        rt = await ctx.runtimes.get(s.session_id)
        pid = product_key(rec)
        assert s.release_id == rid and rt.catalog.order == [pid] and rt.catalog.brand == "B"
        assert rt.catalog.money(rt.catalog.live_price_value(pid)) == "S$36.00"
        assert [p.text for p in ctx.retriever.passages_of(pid) if p.section == "Cautions"] == ["Patch test first."]
        assert rt.pack.rules.rules  # a cosmetic: the market's law pack applies
        shipped = await create_session(ctx, "daon_skincare_demo")  # the shipped catalog is untouched
        assert (await ctx.runtimes.get(shipped.session_id)).catalog.order != [pid]
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()


async def test_d222_the_eval_scripts_run_cases_on_a_release(tmp_path, monkeypatch):
    import sys
    from pathlib import Path

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
    import eval_adversarial

    from app.onboarding.release import save_release

    from .judge_fakes import violation

    rec, _ = normalise(ProductRecord.model_validate(_record()))
    rid = save_release(tmp_path, rec, "job")
    monkeypatch.setattr(eval_adversarial, "RELEASE", rid)
    real_sleep = __import__("asyncio").sleep
    monkeypatch.setattr(__import__("asyncio"), "sleep", lambda _s: real_sleep(0))
    llm = ScriptedBackend()
    llm.judge("It cures eczema.", violations=[violation("ASEAN-2D-DISEASE", "cures eczema")])
    ctx = make_ctx(backend=llm, ONBOARDING_DIR=str(tmp_path))
    case = {"case_id": "R1", "type": "t", "current_product_id": "the product", "utterance": "It cures eczema.",
            "expected": {"risky": True, "kind": "compliance", "rule_id": "ASEAN-2D-DISEASE"}}
    try:
        row = await eval_adversarial.run_case(ctx, case, 0)
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    assert [a["rule_id"] for a in row["alerts"]] == ["ASEAN-2D-DISEASE"]


def test_d223_promo_prices_are_computed_in_code_from_the_stated_terms():
    from decimal import Decimal

    from app.onboarding.release import _rounded, promos_from_offers

    money = {"currency": "SGD", "minor": 6000, "kind": "promo", "market": "SG", "evidence": EV}
    offers = [
        {"offer_id": "o1", "kind": "percent_off", "title": "Extra 15% off", "variant_ids": [], "percent": "15",
         "amount": None, "buy": None, "get": None, "gift": None, "rounding": "rounded to the nearest 10 cents",
         "starts": None, "ends": None, "duration_min": 10, "evidence": EV},
        {"offer_id": "o2", "kind": "spend_threshold_gift", "title": "Free pouch over S$60", "variant_ids": [],
         "percent": None, "amount": money, "buy": None, "get": None, "gift": "pouch", "rounding": None,
         "starts": None, "ends": None, "duration_min": None, "evidence": EV},
        {"offer_id": "o3", "kind": "percent_off", "title": "Other size only", "variant_ids": ["B"], "percent": "50",
         "amount": None, "buy": None, "get": None, "gift": None, "rounding": None, "starts": None, "ends": None,
         "duration_min": None, "evidence": EV},
    ]
    rec = ProductRecord.model_validate(_record(offers=offers))
    sale, gift = promos_from_offers(rec, "p", "A", Decimal("36.00"), "SGD", "SG")  # o3 is for another variant
    assert sale.changes_prices and sale.promo_prices == {"p": 30.6} and sale.duration_sec == 600  # 30.60 to 10 cents
    assert not gift.changes_prices and gift.min_spend == 60.0
    assert _rounded(Decimal("30.66"), None) == Decimal("30.66") and _rounded(Decimal("30.66"), "round down to the dollar") == 30


@pytest.mark.parametrize("written, quote, kept", [
    ("5,000 mg", "Vitamin C 5,000 mg per serving", "5000"),  # D-225: thousands separator, not "5"
    ("5", "Vitamin C 5,000 mg per serving", "5000"),  # D-243: the quote's one number wins, and a person checks it
    ("7", "Vitamin C 5,000 mg and 2 tablets", None),  # not in its quote, and the quote is ambiguous: dropped
    ("0,5", "Vitamin C 0,5 mg", None),  # a decimal comma in an English record is not guessed
])
def test_d225_numbers_must_match_their_evidence(written, quote, kept):
    ev = [{"file": "f", "locator": "l", "quote": quote}]
    rec = ProductRecord.model_validate(_record(ingredients=[{"name": "Vitamin C", "amount_value": written,
                                                             "amount_unit": "mg", "purpose": None, "evidence": ev}]))
    r, notes = normalise(rec)
    assert r.ingredients[0].amount_value == kept
    assert (kept is None) == any("Vitamin C" in m for m in r.missing)


def test_d225_a_decimal_comma_language_reads_commas_as_decimals():
    ev = [{"file": "f", "locator": "l", "quote": "Zink 0,5 mg"}]
    rec = ProductRecord.model_validate(_record(language="de", ingredients=[
        {"name": "Zink", "amount_value": "0,5", "amount_unit": "mg", "purpose": None, "evidence": ev}]))
    assert normalise(rec)[0].ingredients[0].amount_value == "0.5"


async def test_d226_the_caution_pass_adds_what_the_full_read_missed():
    backend = ScriptedBackend()
    backend.outputs[("onboarding_read", "spec.txt")] = _record(cautions=[{"text": "Patch test first.", "evidence": EV}])
    backend.outputs[("onboarding_cautions", "cautions")] = {"cautions": [
        {"text": "patch test first", "evidence": EV},  # already there
        {"text": "Avoid the eye area.", "evidence": EV},
        {"text": "Use sunscreen during the day.", "evidence": []}]}  # no evidence: dropped
    ctx = make_ctx(backend=backend)
    try:
        res = await read_product(ctx.llm, [model_part("spec.txt", b"hello")])
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    assert [c.text for c in res.record.cautions] == ["Patch test first.", "Avoid the eye area."]


def test_d227_ids_in_any_script_are_distinct(tmp_path):
    from app.onboarding.release import product_key, save_release

    a = ProductRecord.model_validate(_record(brand="한빛", product_name="수분 크림"))
    b = ProductRecord.model_validate(_record(brand="한빛", product_name="진정 앰플"))
    assert product_key(a) != product_key(b) and product_key(a).startswith("p_")
    assert save_release(tmp_path, a, "j") != save_release(tmp_path, b, "j")


async def test_d227_a_category_without_a_law_pack_gets_no_law_rules(tmp_path):
    from app.onboarding.release import save_release
    from app.pipelines.runtime import create_session

    rec, _ = normalise(ProductRecord.model_validate(_record(category="apparel")))
    rid = save_release(tmp_path, rec, "j")
    ctx = make_ctx(ONBOARDING_DIR=str(tmp_path))
    try:
        rt = await ctx.runtimes.get((await create_session(ctx, "x", release_id=rid)).session_id)
        assert rt.pack.rules.rules == [] and rt.reflex.scan("This shirt cures eczema.") == []
        assert ctx.pack.rules.rules  # the shipped pack keeps its rules
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()


async def test_d229_an_approved_law_pack_covers_a_category_the_market_pack_does_not(tmp_path):
    from app.onboarding.lawpack import LawPackRecord, normalise_pack, save_lawpack
    from app.onboarding.release import save_release
    from app.pipelines.runtime import create_session

    law, notes = normalise_pack(LawPackRecord.model_validate({
        "name": "Supplement advertising", "regulator": "Example authority", "covers": ["Food Supplement"],
        "language": "en", "rules": [
            {"rule_id": "supp disease", "title": "No disease claims", "description": "No cure or treatment claims.",
             "severity": "HIGH", "rule_category": "Disease", "unacceptable_examples": ["It cures arthritis."],
             "say_instead_hint": "Describe the nutrient's normal function.", "evidence": EV},
            {"rule_id": "X", "title": "t", "description": "d", "severity": "MEDIUM", "rule_category": "c",
             "unacceptable_examples": [], "say_instead_hint": "", "evidence": []}]}))
    assert [r.rule_id for r in law.rules] == ["SUPP-DISEASE"] and law.covers == ["food_supplement"] and notes
    save_lawpack(tmp_path, "SG", law, "j")
    rec, _ = normalise(ProductRecord.model_validate(_record(category="food_supplement")))
    rid = save_release(tmp_path, rec, "j")
    ctx = make_ctx(ONBOARDING_DIR=str(tmp_path))
    try:
        rt = await ctx.runtimes.get((await create_session(ctx, "x", release_id=rid)).session_id)
        assert [r.rule_id for r in rt.pack.rules.rules] == ["SUPP-DISEASE"]
        assert rt.pack.rules_by_id["SUPP-DISEASE"].severity == "HIGH"
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()


def test_d230_upload_approve_start_shows_the_new_product_on_the_session(tmp_path):
    """The user flow end to end: upload, read, approve, start a show; the session's catalog, prompter and run sheet
    are the new product's (not the shipped demo's)."""
    import time

    from fastapi.testclient import TestClient

    from app.main import create_app

    from .helpers import make_settings

    backend = ScriptedBackend()
    backend.outputs[("onboarding_read", "spec.txt")] = _record(
        brand="Lumen", product_name="Night Serum",
        approved_claims=[{"text": "Helps skin look smoother by morning.", "evidence": EV}])
    with TestClient(create_app(make_settings(ONBOARDING_DIR=str(tmp_path)), llm_backend=backend)) as c:
        job = c.post("/api/onboarding/jobs", files=[("files", ("spec.txt", b"hello", "text/plain"))]).json()["job_id"]
        for _ in range(100):
            if c.get(f"/api/onboarding/jobs/{job}").json()["status"] == "done":
                break
            time.sleep(0.02)
        rid = c.post(f"/api/onboarding/jobs/{job}/approve", json={}).json()["release_id"]
        sid = c.post("/api/sessions", json={"release_id": rid}).json()["session_id"]
        cat = c.get(f"/api/catalog?session={sid}").json()
        assert cat["brand"] == "Lumen" and [p["name"] for p in cat["products"]] == ["Night Serum"]
        assert cat["products"][0]["live_price_display"] == "S$36.00"
        assert [rs["title"] for rs in cat["runsheets"]] == ["Lumen · Night Serum"]
        assert c.get("/api/catalog").json()["brand"] != "Lumen"  # the shipped catalog is unchanged
        c.post(f"/api/sessions/{sid}/start")
        state = c.get(f"/api/sessions/{sid}/state?role=host").json()
        # the host prompter is on the new product (its name comes from the session's catalog above)
        assert state["prompter"]["product_id"] == cat["products"][0]["product_id"]
        assert [p["text"] for p in state["prompter"]["points"]] == ["Helps skin look smoother by morning."]


async def test_d232_host_lines_come_in_the_show_language_or_not_at_all(tmp_path):
    from app.onboarding.release import save_release
    from app.pipelines.runtime import create_session

    from .judge_fakes import speak, violation

    rec, _ = normalise(ProductRecord.model_validate(_record(
        language="ms", approved_claims=[{"text": "Membantu kulit kelihatan licin.", "evidence": EV}])))
    with_t = save_release(tmp_path, rec, "j", {"en": {"approved_claims": [{"text": "Helps skin look smooth.",
                                                                          "reviewed": True}],
                                                      "cautions": [], "avoid_saying": []}})
    unreviewed = save_release(tmp_path, rec, "j3", {"en": {"approved_claims": [{"text": "Helps skin look smooth.",
                                                                              "reviewed": False}],
                                                           "cautions": [], "avoid_saying": []}})
    without = save_release(tmp_path, rec, "j2")
    llm = ScriptedBackend()
    llm.judge("It cures eczema.", violations=[violation("ASEAN-2D-DISEASE", "cures eczema", say="Model's own line.")])
    ctx = make_ctx(backend=llm, ONBOARDING_DIR=str(tmp_path))
    try:
        for rid, say in [(with_t, "Helps skin look smooth."),
                         (unreviewed, "Helps skin look smooth. (unreviewed translation)"),  # D-233: marked
                         (without, "(no approved wording in en)")]:
            rt = await ctx.runtimes.get((await create_session(ctx, "x", release_id=rid)).session_id)
            await rt.start()
            await speak(rt, "It cures eczema.")
            assert [a.say_instead for a in await rt.alerts() if a.kind == "compliance"] == [say]
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()


async def test_d233_an_answer_from_an_unreviewed_translation_is_never_auto_posted(tmp_path):
    from app.onboarding.release import UNREVIEWED_DOC, release_pack

    rec, _ = normalise(ProductRecord.model_validate(_record(
        language="ms", cautions=[{"text": "Jangan guna pada kulit luka.", "evidence": EV}],
        qualified_results=[{"statement": "21 of 24 saw smoother skin", "numbers": ["21", "24"],
                            "provenance": "third_party_lab", "qualifier": "tested by Kestrel Labs", "marker": "†",
                            "evidence": EV}],
        approved_claims=[{"text": "Kulit lebih licin dalam 4 minggu†", "evidence": EV}],
        ingredients=[{"name": "Zinc", "amount_value": "0.5", "amount_unit": "mg", "purpose": None, "evidence": EV}])))
    ctx = make_ctx()
    try:
        tr = {"en": {"approved_claims": [{"text": "Smoother skin in 4 weeks†", "reviewed": True}],
                     "cautions": [{"text": "Do not use on broken skin.", "reviewed": False}], "avoid_saying": []}}
        pack = release_pack(ctx.pack, rec, "rel_x_v1", None, tr)
        product = pack.products.products[0]
        assert product.approved_claims == ["Smoother skin in 4 weeks (tested by Kestrel Labs)"]  # marker -> qualifier
        assert product.key_facts == {"zinc": "0.5mg"}  # the prompter's ingredient panel
        [caution] = [p for p in pack.passages if p.section == "Cautions"]
        assert caution.doc == UNREVIEWED_DOC
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()


def test_d236_risk_mode_checks_only_what_reaches_the_show():
    from app.onboarding.checks import check_items, summary
    from app.onboarding.record import OnboardingResult

    rec, _ = normalise(ProductRecord.model_validate(_record(
        cautions=[{"text": "Patch test first.", "evidence": EV}, {"text": "Avoid the eye area.", "evidence": EV}],
        how_to_use=[{"text": "Apply at night.", "evidence": EV}],
        approved_claims=[{"text": "Helps skin feel soft.", "evidence": EV}],
        conflicts=[{"field": "variants[0].prices[0].minor", "values": ["3600", "3400"], "evidence": EV,
                    "effective_dates": [None, None], "note": ""}])))
    result = OnboardingResult(record=rec, method="gemini_whole", seconds=1, caution_pass_added=["Avoid the eye area."])
    full = check_items(result, "full", law="15 rules")
    risk = check_items(result, "risk", law="15 rules", caution_pass_only=result.caution_pass_added)
    assert len(risk) < len(full)
    groups = summary(risk)["by_group"]
    assert groups["product"] == 2 and groups["safety"] == 1 and groups["audit"] == 2  # the conflict, a size with no price
    assert not any(x["label"] == "Apply at night." for x in risk)  # prose is checked live, not here
    assert any(x["reason"].startswith("no law pack") for x in check_items(result, "risk", law="no law checks"))


def test_d238_test_drive_reads_the_sample_approves_and_starts_a_live_show(tmp_path, monkeypatch):
    import time

    from fastapi.testclient import TestClient

    from app.main import create_app
    from app.pipelines import testdrive as td

    from .helpers import make_settings

    monkeypatch.setattr(td, "_starts", [])
    monkeypatch.setattr(td, "_admitted", {})
    monkeypatch.setattr(td, "_running", set())
    monkeypatch.setattr(td, "MAX_PER_HOUR", 1)
    played = []

    async def no_scene(rt):
        played.append(rt.sid)

    monkeypatch.setattr(td, "play", no_scene)
    backend = ScriptedBackend()
    settings = make_settings(ONBOARDING_DIR=str(tmp_path))
    names = "|".join(n for n, _ in td.sample_files(settings.data_path))
    backend.outputs[("onboarding_read", names)] = _record(brand="Lumen", product_name="Night Serum")
    with TestClient(create_app(settings, llm_backend=backend)) as c:
        start = c.post("/api/testdrive/start").json()
        assert not start["cached"] and len(start["files"]) == 4
        for _ in range(100):
            if c.get(f"/api/onboarding/jobs/{start['job_id']}").json()["status"] == "done":
                break
            time.sleep(0.02)
        assert c.post("/api/testdrive/job_other/go").status_code == 429  # D-322: only the admitted job
        go = c.post(f"/api/testdrive/{start['job_id']}/go").json()
        state = c.get(f"/api/sessions/{go['session_id']}/state?role=host").json()
        assert state["session"]["status"] == "live" and played == [go["session_id"]]
        assert c.post(f"/api/testdrive/{start['job_id']}/go").status_code == 429  # D-322: once
        assert c.post("/api/testdrive/start").status_code == 429  # the hourly cap
        monkeypatch.setattr(td, "_starts", [])
        again = c.post("/api/testdrive/start").json()
        assert again["cached"] and again["status"] == "done"  # the sample is read once, then reused


async def test_d240_impact_counts_come_with_their_denominators():
    from .helpers import live_runtime, say
    from .judge_fakes import speak, violation

    llm = ScriptedBackend()
    ctx = make_ctx(backend=llm)
    rt = await live_runtime(ctx)
    try:
        llm.judge("It cures eczema.", violations=[violation("ASEAN-2D-DISEASE", "cures eczema")])
        await speak(rt, "It cures eczema.")
        await speak(rt, "It feels light.")
        await say(rt, "Lee", "How much is the cream?")
        i = (await rt.compute_metrics())["impact"]
        assert i["risky_lines"] == {"alerted": 1, "of_lines": 2}
        assert i["questions"]["of_askers"] == 1 and i["questions"]["ai"] + i["questions"]["host"] + i["questions"]["operator"] >= 1
        assert i["operator_actions"] == 0 and i["llm"]["calls"] > 0
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()


async def test_d241_brand_rules_only_add_restrictions(tmp_path):
    from app.onboarding.lawpack import LawPackRecord, normalise_pack, save_brandpack
    from app.onboarding.release import release_pack

    brand, _ = normalise_pack(LawPackRecord.model_validate({
        "name": "B guidelines", "regulator": "B", "covers": ["cosmetic"], "language": "en", "rules": [
            {"rule_id": "NO-CELEB", "title": "No celebrity comparisons", "description": "Never compare to a celebrity.",
             "severity": "MEDIUM", "rule_category": "comparison", "unacceptable_examples": ["Skin like a movie star."],
             "say_instead_hint": "Describe the product itself.", "evidence": EV}]}))
    save_brandpack(tmp_path, "B", brand, "j")
    rec, _ = normalise(ProductRecord.model_validate(_record()))  # brand "B", a cosmetic
    ctx = make_ctx()
    try:
        law = release_pack(ctx.pack, rec, "rel_x_v1", None).rules.rules
        both = release_pack(ctx.pack, rec, "rel_x_v1", tmp_path).rules.rules
        assert [r.rule_id for r in both] == [r.rule_id for r in law] + ["BR-NO-CELEB"]  # law kept, brand added
        other = rec.model_copy(update={"brand": "Other"})
        assert len(release_pack(ctx.pack, other, "rel_y_v1", tmp_path).rules.rules) == len(law)
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()


def test_d242_record_mode_gives_the_fact_checker_every_size_market_offer_and_superseded_value():
    from app.onboarding.release import release_pack

    ev = EV
    rec, _ = normalise(ProductRecord.model_validate(_record(
        variants=[{"variant_id": "A", "name": "P 30", "size_value": "30", "size_unit": "ml", "market": None,
                   "prices": [{"currency": "SGD", "minor": 3600, "kind": "live", "market": "SG", "evidence": ev},
                              {"currency": "MYR", "minor": 11900, "kind": "live", "market": "MY",
                               "evidence": [{"file": "f", "locator": "l", "quote": "RM119"}]}], "evidence": ev}],
        offers=[{"offer_id": "o1", "kind": "buy_x_get_y", "title": "Buy 2 get 1", "variant_ids": [], "percent": None,
                 "amount": None, "buy": 2, "get": 1, "gift": None, "rounding": None, "starts": None, "ends": None,
                 "duration_min": None, "evidence": ev}],
        conflicts=[{"field": "P 30 live price", "values": ["S$38.00", "S$36.00"],
                    "effective_dates": ["2026-09-01", "2026-10-01"], "note": "", "evidence": ev}])))
    ctx = make_ctx()
    try:
        sheet = release_pack(ctx.pack, rec, "rel_x_v1", None, fact_mode="sheet")
        record = release_pack(ctx.pack, rec, "rel_x_v1", None, fact_mode="record")
        assert not [p for p in sheet.passages if p.section == "Record facts"]
        facts = [p.text for p in record.passages if p.section == "Record facts"]
        assert "P 30 (30 ml) · MY live price: MYR 119.00" in facts and "P 30 (30 ml) · SG live price: S$36.00" in facts
        assert any(t.startswith("Superseded: P 30 live price was S$38.00") for t in facts)
        assert any(t.startswith("Offer: Buy 2 get 1") and "buy 2 get 1" in t for t in facts)
        assert "Live broadcast price (30 ml): S$36.00" in [p.text for p in record.passages]
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()


@pytest.mark.parametrize("quote, written, currency, kept_minor", [
    ("가격 이만 구천 원", "290", "KRW", 29000),  # D-243: Korean words; the model's misread is replaced
    ("Harga dua puluh sembilan ringgit", "29", "MYR", 2900),  # Malay words
    ("售价二万九千", "29000", "KRW", 29000),  # Chinese numerals
    ("Now only 2만 9천원", "29000", "KRW", 29000),  # mixed digits and Korean scale words
    ("Price: twenty-nine dollars", "29", "SGD", 2900),
])
def test_d243_prices_written_in_words_are_read_in_code(quote, written, currency, kept_minor):
    from app.onboarding.record import to_minor

    ev = [{"file": "f", "locator": "l", "quote": quote}]
    rec = ProductRecord.model_validate(_record(variants=[{
        "variant_id": "A", "name": "P", "size_value": None, "size_unit": None, "market": None,
        "prices": [{"currency": currency, "minor": to_minor(currency, written), "kind": "list", "market": None,
                    "evidence": ev}], "evidence": ev}]))
    r, notes = normalise(rec)
    assert r.variants[0].prices[0].minor == kept_minor


def test_d243_zero_decimal_currencies_keep_the_whole_amount():
    from app.onboarding.record import money_value, to_minor

    assert money_value("KRW", 29000) == 29000 and to_minor("KRW", "29000") == 29000
    assert money_value("SGD", 3600) == 36 and to_minor("SGD", "36") == 3600


async def test_d243_routing_turns_tests_and_certifications_into_qualified_results_and_limitations_into_cautions():
    backend = ScriptedBackend()
    backend.outputs[("onboarding_read", "spec.txt")] = _record(
        key_facts=[{"text": "Tested by Kestrel Labs: 21 of 24 felt softer skin.", "evidence": EV}],
        approved_claims=[{"text": "Certified vegan by the Vegan Society.", "evidence": EV}],
        how_to_use=[{"text": "Results may vary between users.", "evidence": EV},
                    {"text": "Keep the applicator clean.", "evidence": EV}])
    nolink = {"linked": None, "footnote_quote": None, "footnote_file": None}
    backend.outputs[("onboarding_route", "route|P")] = {"items": [x | nolink for x in [
        {"id": "key_facts[0]", "kind": "test_result", "statement": "21 of 24 felt softer skin",
         "provenance": "third_party_lab", "qualifier": "tested by Kestrel Labs", "marker": None},
        {"id": "approved_claims[0]", "kind": "certification", "statement": "Certified vegan",
         "provenance": "other", "qualifier": "certified by the Vegan Society", "marker": None},
        {"id": "how_to_use[0]", "kind": "limitation", "statement": None, "provenance": None, "qualifier": None,
         "marker": None},
        {"id": "nope[9]", "kind": "test_result", "statement": "x", "provenance": "other", "qualifier": "q",
         "marker": None}]]}
    # D-243 F3: the caution pass lists a hygiene line that also stays in the directions
    backend.outputs[("onboarding_cautions", "cautions")] = {"cautions": [{"text": "Keep the applicator clean.",
                                                                           "evidence": EV}]}
    ctx = make_ctx(backend=backend)
    try:
        res = await read_product(ctx.llm, [model_part("spec.txt", b"hello")])
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    r = res.record
    assert [(q.statement, q.provenance, q.qualifier) for q in r.qualified_results] == [
        ("21 of 24 felt softer skin", "third_party_lab", "tested by Kestrel Labs"),
        ("Certified vegan", "other", "certified by the Vegan Society")]
    assert r.qualified_results[0].numbers == ["21", "24"] and r.qualified_results[0].evidence == r.key_facts[0].evidence
    assert [c.text for c in r.cautions] == ["Keep the applicator clean.", "Results may vary between users."]
    assert "Keep the applicator clean." in [h.text for h in r.how_to_use]  # kept in the directions too


def test_d244_l1_a_contradiction_without_a_source_value_is_unverifiable(pack, catalog):
    from app.agents.fact_checker import post_validate
    from app.models.agent_io import FactCheckOutput

    cream = "cream_hydra_barrier"
    quote = "Live broadcast price: S$36.00"
    pid = "DL-CR-01#key-facts#live-broadcast-price"
    claim = {"span": "x", "field": "other", "product_id": cream, "stated_value": "5 mg vitamin E",
             "verdict": "contradicted", "correct_value": "not disclosed", "passage_id": pid, "quote": quote}
    out = FactCheckOutput.model_validate({"claims": [claim | {"field": "ingredient_amount"}], "reason": "t"})
    assert post_validate(out, pack.passages, catalog, None).claims[0].verdict == "not_in_sources"


def test_d244_l2_l3_offer_totals_and_written_terms_reach_the_fact_checker():
    from app.onboarding.release import release_pack

    ev = [{"file": "f", "locator": "l", "quote": "Buy 2 get 1 free; not combinable with other offers"}]
    rec, _ = normalise(ProductRecord.model_validate(_record(offers=[
        {"offer_id": "o1", "kind": "buy_x_get_y", "title": "Buy 2 get 1", "variant_ids": [], "percent": None,
         "amount": None, "buy": 2, "get": 1, "gift": None, "rounding": None, "starts": None, "ends": None,
         "duration_min": None, "evidence": ev}])))
    ctx = make_ctx()
    try:
        facts = [p.text for p in release_pack(ctx.pack, rec, "rel_x_v1", None).passages if p.section == "Record facts"]
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    [offer] = [t for t in facts if t.startswith("Offer:")]
    assert "pay for 2 = S$72.00 for 3 items" in offer and "not combinable with other offers" in offer
    assert "Live unit price S$36.00; 2 units S$72.00; 3 units S$108.00" in facts


def test_d244_l7_naming_the_lab_is_saying_the_qualifier(pack):
    from app.agents.utterance_judge import names_the_source

    say = "21 of 24 saw smoother texture (tested by Kestrel Dermal Labs)"
    assert names_the_source("Kestrel Dermal Labs found 21 of 24 saw smoother texture", say, pack.lexicon)
    assert not names_the_source("21 of 24 saw smoother texture", say, pack.lexicon)
    assert not names_the_source("our test: 45% more", "Hydration +45% (in our in-house test)", pack.lexicon)


def test_d244_l4_to_l6_and_presence_rules_reach_the_fact_checker():
    from app.agents import arms
    from app.config import get_settings

    fr = arms._addenda(str(get_settings().prompts_path))[("FR", "fact_checker")]
    for rule in ("never with\n  the unit price", "Every number with a unit", "Price terms", "describing an old value",
                 "Free from", "Only contradict with a value SOURCE states"):
        assert rule in fr


def test_d244_l8_a_bare_field_word_becomes_the_source_line(pack, catalog):
    from app.agents.fact_checker import _from_quote, post_validate
    from app.models.agent_io import FactCheckOutput

    p = next(p for p in pack.passages if p.product_id == "cream_hydra_barrier")
    quote = p.text[:30]
    word = quote.split()[0]
    claim = {"span": f"no {word} in it", "field": "ingredient", "product_id": "cream_hydra_barrier",
             "stated_value": f"no {word}", "verdict": "contradicted", "correct_value": word,
             "passage_id": p.passage_id, "quote": quote}
    [c] = post_validate(FactCheckOutput.model_validate({"claims": [claim], "reason": "t"}),
                        pack.passages, catalog, None).claims
    assert c.verdict == "contradicted" and c.correct_value == quote
    # a statement in the quote's words passes; "not disclosed" does not, in any script
    assert _from_quote("It does contain fragrance", "Ingredients: Water, Glycerin, Fragrance")
    assert not _from_quote("not disclosed", "Live broadcast price: S$36.00")
    assert _from_quote("향료 함유", "전성분: 정제수, 글리세린, 향료")
    assert not _from_quote("공개되지 않음", "라이브 가격: 36,000원")


async def test_d245_r1_a_marked_claim_and_its_footnote_become_one_qualified_result():
    backend = ScriptedBackend()
    claim_ev = [{"file": "spec.txt", "locator": "p1", "quote": "UV shield SPF40 †"}]
    note_ev = [{"file": "spec.txt", "locator": "p2", "quote": "† Tested by Larkspur Labs, report LL-77"}]
    backend.outputs[("onboarding_read", "spec.txt")] = _record(
        approved_claims=[{"text": "UV shield SPF40 †", "evidence": claim_ev}],
        key_facts=[{"text": "† Tested by Larkspur Labs, report LL-77", "evidence": note_ev}])
    backend.outputs[("onboarding_route", "route|P")] = {"items": [
        {"id": "approved_claims[0]", "kind": "test_result", "statement": "UV shield SPF40",
         "provenance": "third_party_lab", "qualifier": "tested by Larkspur Labs", "marker": "†",
         "linked": "key_facts[0]", "footnote_quote": "† Tested by Larkspur Labs, report LL-77",
         "footnote_file": "spec.txt"},
        {"id": "key_facts[0]", "kind": "none", "statement": None, "provenance": None, "qualifier": None,
         "marker": None, "linked": None, "footnote_quote": None, "footnote_file": None}]}
    ctx = make_ctx(backend=backend)
    try:
        res = await read_product(ctx.llm, [model_part("spec.txt", b"hello")])
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    [q] = res.record.qualified_results
    assert (q.provenance, q.qualifier, q.marker, q.numbers) == ("third_party_lab", "tested by Larkspur Labs", "†", ["40"])
    quotes = [e.quote for e in q.evidence]
    assert quotes[0] == "UV shield SPF40 †" and "† Tested by Larkspur Labs, report LL-77" in quotes
    assert [e.locator for e in q.evidence][-1] == "footnote"


async def test_d245_r2_a_free_text_category_is_mapped_to_the_closed_list(tmp_path):
    from app.onboarding.read import category_list

    ctx = make_ctx()
    try:
        cats = category_list(ctx.pack, tmp_path)
        assert ctx.pack.market.category in cats and "apparel" in cats and len(cats) == len(set(cats))
        backend = ScriptedBackend()
        backend.outputs[("onboarding_read", "spec.txt")] = _record(category="sports_leggings")
        backend.outputs[("onboarding_category", "category|sports_leggings")] = {"category": "Apparel"}
        ctx2 = make_ctx(backend=backend)
        try:
            res = await read_product(ctx2.llm, [model_part("spec.txt", b"hello")], categories=cats)
        finally:
            ctx2.runtimes.shutdown()
            ctx2.scheduler.shutdown()
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    r = res.record
    assert r.category == "apparel" and r.category_evidence[-1].quote == "sports_leggings"
    assert "category: read as sports_leggings, mapped to apparel" in res.notes


def test_d245_r3_a_certification_body_is_a_certifier():
    from app.onboarding.record import QualifiedResult

    q = QualifiedResult.model_validate({"statement": "Certified vegan", "provenance": "certifier",
                                        "qualifier": "certified by the Vegan Society", "evidence": []})
    assert q.provenance == "certifier"


@pytest.mark.parametrize("text, values", [
    ("영 점 오 퍼센트", {"0.5"}), ("영점오", {"0.5"}), ("일 점 이 ml", {"1.2"}), ("零点五", {"0.5"}),
    ("2만 9천 원", {"29000"}), ("1.5만", {"15000"}),  # the digits are part of the numeral, not numbers of their own
    ("5 이하", {"5"}), ("단점이 없다", set()),  # a digit then a word, a word with 점
])
def test_d245_r4_korean_and_chinese_decimals(text, values):
    from app.onboarding.numwords import numbers_in_text, plain

    assert {plain(v) for v in numbers_in_text(text)} == values


def test_d245_r4_r5_a_word_decimal_is_kept_and_a_dropped_amount_reaches_the_check_list():
    from app.onboarding.checks import check_items
    from app.onboarding.record import OnboardingResult

    ing = [{"name": "Zinc", "amount_value": "0.5", "amount_unit": "%", "purpose": None,
            "evidence": [{"file": "f", "locator": "l", "quote": "아연 영 점 오 퍼센트"}]},
           {"name": "Panthenol", "amount_value": "2", "amount_unit": "%", "purpose": None,
            "evidence": [{"file": "f", "locator": "l", "quote": "판테놀 함유"}]}]
    rec, notes = normalise(ProductRecord.model_validate(_record(ingredients=ing)))
    assert [i.amount_value for i in rec.ingredients] == ["0.5", None]
    risk = check_items(OnboardingResult(record=rec, method="gemini_whole", seconds=1, notes=notes), "risk", law="15 rules")
    assert [x["group"] for x in risk if x["label"] == "ingredients[Panthenol].amount"] == ["product"]


def test_d246_facts_model_is_the_fact_checkers_own():
    from app.agents.fact_checker import FactChecker

    for kw, want in (({}, "gemini-3.5-flash-lite"), ({"FACTS_MODEL": "gemini-3.8-flash"}, "gemini-3.8-flash")):
        ctx = make_ctx(LLM_FAST_MODEL="gemini-3.5-flash-lite", **kw)
        try:
            assert ctx.llm.facts.model == want and ctx.llm.fast.model == "gemini-3.5-flash-lite"
            assert FactChecker(ctx.llm, ctx.prompts["fact_checker"], None).model == want
        finally:
            ctx.runtimes.shutdown()
            ctx.scheduler.shutdown()


async def test_d247_fd_the_fact_checker_reads_the_release_files_and_code_checks_the_quote(tmp_path):
    from app.agents.fact_checker import FactChecker, FactInput
    from app.logic.catalog import Catalog
    from app.onboarding.files import store
    from app.onboarding.release import save_release
    from app.models.session import SwitchState
    from app.onboarding.sources import release_sources

    rec, _ = normalise(ProductRecord.model_validate(_record()))
    spec = b"Contains a light citrus fragrance.\nUse within 12 months of opening."
    store(tmp_path, "job1", "spec.txt", spec)
    store(tmp_path, "job1", "label.png", b"\x89PNG not really")
    rid = save_release(tmp_path, rec, "job1")
    parts, passages = release_sources(str(tmp_path), rid)
    assert [m for _, m in parts] == ["text/plain", "text/plain", "text/plain", "image/png"]
    by_id = {p.passage_id: p for p in passages}
    assert "12 months of opening" in by_id["file:spec.txt"].text and by_id["file:label.png"].section == "Source file"

    backend = ScriptedBackend()
    good = {"span": "use it within 24 months", "field": "shelf_life", "stated_value": "24 months",
            "verdict": "contradicted", "correct_value": "12 months", "passage_id": "file:spec.txt#line 2",
            "quote": "Use within 12 months of opening."}
    made_up = good | {"quote": "Use within 12 months, always.", "span": "within 24 months, no rush"}
    ctx = make_ctx(backend=backend, ONBOARDING_DIR=str(tmp_path), FACT_MODE="docs")
    try:
        pack, _ = ctx.pack_for(ctx.pack.market.market, ctx.pack.lang, rid)
        catalog = Catalog(pack, rid)
        pid = catalog.order[0]
        fc = FactChecker(ctx.llm, ctx.prompts["fact_checker"], catalog)
        assert fc.file_parts == parts and "Fact mode \"docs\"" in fc.system_instruction
        assert "Fact mode \"record\"" in fc.system_instruction  # FD is FR plus the files
        out = []
        for claim in (good, made_up):
            backend.facts(claim["span"], claim | {"product_id": pid})
            res = await fc.run(FactInput(utterance=claim["span"], on_air=pid, switch=SwitchState(), promo=None, passages=[],
                                         session_id="s1"), deadline_ms=3000)
            out.append((res.output.claims[0].verdict, res.output.claims[0].passage_id))
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    assert out == [("contradicted", "file:spec.txt"), ("not_in_sources", "file:spec.txt")]


def test_d247_fd_the_files_are_cached_per_release_with_an_empty_prefix():
    from app.models.agent_io import FactCheckOutput
    from app.providers.llm import LLMRequest, VertexBackend

    vb = VertexBackend.__new__(VertexBackend)
    vb._context_cache, vb._caches = True, {}
    req = lambda files, prefix="": LLMRequest(  # noqa: E731
        agent="fact_checker", model="gemini-3.5-flash-lite", thinking_level="MINIMAL", prompt_version="v",
        system_instruction="s", user_message=prefix + "line", output_model=FactCheckOutput, deadline_ms=1000,
        session_id="s1", cache_prefix=prefix, cache_parts=files)
    a, b = ((b"FILE a", "text/plain"), (b"one", "text/plain")), ((b"FILE a", "text/plain"), (b"two", "text/plain"))
    assert vb._cache_entry(req(a)) is vb._cache_entry(req(a)) and vb._cache_entry(req(a)) is not vb._cache_entry(req(b))
    assert vb._cache_entry(req(())) is None  # no files and no prefix: nothing to cache
    contents = VertexBackend._contents(req(a), "line")
    assert len(contents) == 3 and contents[-1] == "line"


async def _release_line(tmp_path, monkeypatch, record: dict, line: str, claim_of, row_out: dict | None = None,
                        fail: tuple[str, ...] = ()) -> list[dict]:
    """D-248: one host line through the real pipeline (eval_adversarial.run_case) on a made-up release; the fact
    checker's output is scripted from the release's own passages by `claim_of(pid, passages)`."""
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
    import eval_adversarial

    from app.onboarding.release import product_key, release_pack, save_release

    real_sleep = __import__("asyncio").sleep
    monkeypatch.setattr(__import__("asyncio"), "sleep", lambda _s: real_sleep(0))
    rec, _ = normalise(ProductRecord.model_validate(_record(**record)))
    rid = save_release(tmp_path, rec, "job")
    monkeypatch.setattr(eval_adversarial, "RELEASE", rid)
    backend = ScriptedBackend()
    ctx = make_ctx(backend=backend, ONBOARDING_DIR=str(tmp_path), REFLEX_ON_FAIL="alert")  # D-293: the alert path
    try:
        pack = release_pack(ctx.pack, rec, rid, tmp_path)
        pid = product_key(rec)
        backend.facts(line, claim_of(pid, [p for p in pack.passages if p.product_id == pid]))
        for agent in fail:  # these agents fail on the line
            backend.fail(agent, line)
        case = {"case_id": "X", "type": "t", "current_product_id": pid, "utterance": line,
                "expected": {"risky": True, "kind": "fact"}}
        row = await eval_adversarial.run_case(ctx, case, 0)
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    if row_out is not None:
        row_out.update(row)
    return [a for a in row["alerts"] if a["kind"] == "fact"]


def _claim(pid, p, **kw) -> dict:
    return {"span": "x", "product_id": pid, "verdict": "contradicted", "passage_id": p.passage_id, "quote": p.text} | kw


async def test_d248_a_stated_number_corrected_by_no_value_never_reaches_the_host(tmp_path, monkeypatch):
    def claim(pid, ps):
        p = next(p for p in ps if "Patch test" in p.text)
        return _claim(pid, p, field="shelf_life", stated_value="24 months",
                      correct_value="No period after opening is given")

    row: dict = {}
    alerts = await _release_line(tmp_path, monkeypatch, {"cautions": [{"text": "Patch test first.", "evidence": EV}]},
                                 "It lasts twenty four months after opening.", claim, row)
    assert alerts == []
    # D-250: the eval row says which code check overruled the model
    assert [d["reason"] for d in row["dropped"]] == ["L1: the correction states no value of the quote"]


async def test_d248_ingredient_presence_both_ways_reaches_the_host(tmp_path, monkeypatch):
    facts = [{"text": "No fragrance ingredient is added.", "evidence": EV}]

    def added(pid, ps):
        p = next(p for p in ps if p.text == "No fragrance ingredient is added.")
        return _claim(pid, p, field="fragrance", stated_value="contains fragrance",
                      correct_value="No fragrance ingredient is added.")

    [a] = await _release_line(tmp_path / "a", monkeypatch, {"key_facts": facts}, "There's added perfume in it.", added)
    assert a["correct_value"] == "No fragrance ingredient is added."

    def absent(pid, ps):
        p = next(p for p in ps if p.text == "Contains a light citrus fragrance.")
        return _claim(pid, p, field="fragrance", stated_value="contains no fragrance", correct_value="contains fragrance")

    [b] = await _release_line(tmp_path / "b", monkeypatch,
                              {"key_facts": [{"text": "Contains a light citrus fragrance.", "evidence": EV}]},
                              "There's no fragrance inside.", absent)
    assert b["field"] == "fragrance"


async def test_d248_a_threshold_is_not_the_number_alone(tmp_path, monkeypatch):
    ev = [{"file": "f", "locator": "l", "quote": "Free pouch when the subtotal is more than S$90"}]
    offer = {"offer_id": "o1", "kind": "spend_threshold_gift", "title": "Free pouch", "variant_ids": [], "percent": None,
             "amount": {"currency": "SGD", "minor": 9000, "kind": "promo", "market": "SG", "evidence": ev},
             "buy": None, "get": None, "gift": "pouch", "rounding": None, "starts": None, "ends": None,
             "duration_min": None, "evidence": ev}

    def claim(pid, ps):
        p = next(p for p in ps if p.text.startswith("Offer: Free pouch"))
        assert "minimum spend S$90.00" in p.text and "gift: pouch" in p.text
        return _claim(pid, p, field="promo", stated_value="S$90", correct_value="S$90",
                      quote="Free pouch when the subtotal is more than S$90")

    [a] = await _release_line(tmp_path, monkeypatch, {"offers": [offer]}, "Exactly ninety dollars qualifies.", claim)
    assert a["field"] == "promo" and a["correct_value"] == "Free pouch when the subtotal is more than S$90"


def test_d248_the_qualifier_rule_comes_with_the_records_results():
    from app.logic.qualifier import RULE_ID
    from app.onboarding.release import has_law, release_pack

    result = [{"statement": "Hydration up 18% after 8 hours", "numbers": ["18", "8"], "provenance": "in_house",
               "qualifier": "in our in-house test", "marker": None, "evidence": EV}]
    rec, _ = normalise(ProductRecord.model_validate(_record(category="garden_tools", qualified_results=result)))
    ctx = make_ctx()
    try:
        rules = release_pack(ctx.pack, rec, "rel_x_v1", None).rules
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    assert [r.rule_id for r in rules.rules] == [RULE_ID] and not has_law(rules)


def test_d248_a_sentence_that_gives_no_value_is_no_correction():
    from app.agents.fact_checker import _from_quote

    assert not _from_quote("No water-resistance claim is made", "Long-wear formula for daily use")
    assert not _from_quote("No expiry date or period after opening is given", "Use after opening as directed.")
    assert _from_quote("It does contain fragrance", "Ingredients: Water, Glycerin, Fragrance")


async def test_d249_the_coverage_audit_lists_what_the_record_lacks_and_checks_its_quote():
    from app.onboarding.checks import check_items

    backend = ScriptedBackend()
    backend.outputs[("onboarding_read", "spec.txt|label.jpg")] = _record()
    backend.outputs[("onboarding_audit", "audit|P")] = {"items": [
        {"kind": "condition", "text": "pouch only over S$60", "file": "spec.txt", "locator": "line 2",
         "quote": "Free pouch on orders over S$60", "record_has": None},
        {"kind": "number", "text": "use within 6 months", "file": "spec.txt", "locator": "line 3",
         "quote": "Use within 9 months", "record_has": None},
        {"kind": "claim", "text": "dermatologist tested", "file": "label.jpg", "locator": "image",
         "quote": "Dermatologist tested", "record_has": None},
        {"kind": "condition", "text": "dup", "file": "spec.txt", "locator": "line 2",
         "quote": "Free pouch on orders  over S$60", "record_has": None}]}
    ctx = make_ctx(backend=backend)
    try:
        res = await read_product(ctx.llm, [model_part("spec.txt", b"P 30 ml\nFree pouch on orders over S$60\n"
                                                                  b"Use within 6 months"),
                                           model_part("label.jpg", b"\xff\xd8")])
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    assert [(x["text"], x["in_file_text"]) for x in res.uncovered] == [
        ("pouch only over S$60", True), ("use within 6 months", False), ("dermatologist tested", None)]
    items = [x for x in check_items(res, "risk", law="15 rules") if x["group"] == "coverage"]
    assert len(items) == 3 and "spec.txt (line 2)" in items[0]["reason"] and "not found" in items[1]["reason"]
    assert len([x for x in check_items(res, "full", law="15 rules") if x["group"] == "coverage"]) == 3


async def test_d250_a_correction_in_number_words_is_a_number(tmp_path, monkeypatch):
    def claim(pid, ps):
        p = next(p for p in ps if p.text == "Use within twelve months after opening.")
        return _claim(pid, p, field="shelf_life", stated_value="24 months", correct_value="twelve months after opening")

    [a] = await _release_line(tmp_path, monkeypatch,
                              {"how_to_use": [{"text": "Use within twelve months after opening.", "evidence": EV}]},
                              "Use it within twenty-four months after opening.", claim)
    assert a["correct_value"] == "twelve months after opening"


def test_d250_the_panel_or_the_hours_never_name_an_in_house_source(pack):
    from app.agents.utterance_judge import names_the_source

    line = "hydration was 18% above baseline eight hours after one application in 22 adults"
    lex = pack.lexicon
    assert not names_the_source(line, "Hydration up 18% after 8 hours (in our in-house test with 22 adults)", lex)
    assert not names_the_source(line, "Hydration up 18% after 8 hours (consumer test with 22 adults)", lex)
    assert names_the_source("Kestrel Dermal Labs found 21 of 24", "21 of 24 (tested by Kestrel Dermal Labs)", lex)


async def test_d251_a_quote_with_the_number_in_korean_words_backs_a_digit_correction(tmp_path, monkeypatch):
    text = "나이아신아마이드가 영 점 오 퍼센트 들어 있습니다."

    def claim(pid, ps):
        p = next(p for p in ps if p.text == text)
        return _claim(pid, p, field="ingredient_amount", stated_value="2% niacinamide", correct_value="0.5%")

    [a] = await _release_line(tmp_path, monkeypatch, {"key_facts": [{"text": text, "evidence": EV}]},
                              "It has two percent niacinamide.", claim)
    assert a["correct_value"] == "0.5%"


async def test_d251_a_right_price_with_a_wrong_term_keeps_its_alert(tmp_path, monkeypatch):
    facts = [{"text": "Delivery is not included in this price.", "evidence": EV},
             {"text": "No period after opening is printed.", "evidence": EV}]

    def term(pid, ps):
        p = next(p for p in ps if p.text == "Delivery is not included in this price.")
        return _claim(pid, p, field="price", stated_value="S$36 per tube including delivery",
                      correct_value="Delivery is not included in this price.")

    [a] = await _release_line(tmp_path / "a", monkeypatch, {"key_facts": facts},
                              "Thirty-six dollars per tube already includes delivery.", term)
    assert a["field"] == "price"

    def period(pid, ps):  # the number itself has no source: still dropped
        p = next(p for p in ps if p.text == "No period after opening is printed.")
        return _claim(pid, p, field="shelf_life", stated_value="9 months",
                      correct_value="No period after opening is printed.")

    row: dict = {}
    assert await _release_line(tmp_path / "b", monkeypatch, {"key_facts": facts},
                               "Use it within nine months after opening.", period, row) == []
    assert [d["reason"] for d in row["dropped"]] == ["L1: a stated number corrected without a number"]


@pytest.mark.parametrize("line, shown", [
    ("A viewer says, 'it cures eczema,' and yes, that is what the cream does.", True),
    ("A viewer calls it the best in the country, and I fully agree.", True),
    ("A viewer says it cures eczema, but I don't agree, it is a cosmetic.", False),
])
async def test_d252_a_quote_the_host_endorses_is_the_hosts_own_claim(line, shown):
    from .judge_fakes import close_env, make_env, speak, violation

    rt, llm = await make_env()
    try:
        llm.judge(line, violations=[violation("ASEAN-2D-DISEASE", "cures eczema", quoted=True)])
        await speak(rt, line)
        active = [a for a in await rt.alerts() if a.kind == "compliance" and a.status == "active"]
    finally:
        await close_env(rt)
    assert bool(active) is shown


def test_d252_offer_combinability_and_conditions_are_labelled():
    from app.onboarding.release import release_pack

    ev = [{"file": "f", "locator": "l", "quote": "10% off; cannot be combined with other codes"}]
    rec, _ = normalise(ProductRecord.model_validate(_record(offers=[
        {"offer_id": "o1", "kind": "percent_off", "title": "10% off", "variant_ids": [], "percent": "10", "amount": None,
         "buy": None, "get": None, "gift": None, "rounding": None, "starts": None, "ends": None, "duration_min": None,
         "combinable": False, "conditions": ["one per order"], "evidence": ev}])))
    ctx = make_ctx()
    try:
        facts = [p.text for p in release_pack(ctx.pack, rec, "rel_x_v1", None).passages if p.text.startswith("Offer:")]
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    assert "combinable with other offers or codes: no" in facts[0] and "condition: one per order" in facts[0]


async def test_d252_an_absence_the_quote_does_not_state_is_no_correction(tmp_path, monkeypatch):
    facts = [{"text": "Reapply after swimming or water exposure and every 2 hours.", "evidence": EV}]

    def claim(pid, ps):
        p = next(p for p in ps if p.text == facts[0]["text"])
        return _claim(pid, p, field="usage", stated_value="reapply after swimming",
                      correct_value="No water-resistance claim is made",
                      quote="Reapply after swimming or water exposure")

    row: dict = {}
    assert await _release_line(tmp_path, monkeypatch, {"key_facts": facts},
                               "Reapply after swimming, and also every two hours.", claim, row) == []
    assert [d["reason"] for d in row["dropped"]] == ["L1: an absence the quote does not state"]


def test_d253_a_file_quote_survives_pdf_table_text():
    from app.logic.verify import quote_in_file

    pdf = ("[page 1] Protection | SPF 50 +\nPA + + + +\n│ Water\nresistance │ 80 min-\nutes †\n"
           "Tested by Larkspur\nLabs, report LL-77 ─────")
    assert quote_in_file("SPF50+ PA++++", pdf)
    assert quote_in_file("Water resistance 80 minutes †", pdf)  # a line break inside a word
    assert quote_in_file("Tested by Larkspur Labs, report LL-77", pdf)
    assert not quote_in_file("Water resistance 40 minutes", pdf)  # every number of the quote must be there
    assert not quote_in_file("Clinically proven to cure acne", pdf)


def test_d253_fd_sends_the_qualified_results_to_the_compliance_judge():
    from app.agents.split_judges import ComplianceJudge

    from app.logic.catalog import Catalog

    ctx = make_ctx(FACT_MODE="docs")
    try:
        judge = ComplianceJudge(ctx.llm, ctx.prompts, Catalog(ctx.pack), ctx.reflex)
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    assert judge.fact_record and "QUALIFIED_RESULTS" in judge.system_instruction + judge.extra_data()


def test_d254_lean_checks_at_most_ten_and_folds_or_accepts_the_rest():
    from app.onboarding.checks import LEAN_MAX, check_items, lean_items
    from app.onboarding.record import OnboardingResult

    spec = "Keep away from eyes. Store below 30C. Do not use on broken skin. Stop if irritation occurs."
    ev = lambda q, f="spec.txt": [{"file": f, "locator": "p1", "quote": q}]  # noqa: E731
    cautions = [{"text": t, "evidence": ev(t)} for t in
                ("Keep away from eyes.", "Store below 30C.", "Do not use on broken skin.", "Stop if irritation occurs.")]
    cautions.append({"text": "Not for children.", "evidence": ev("Not for children.", "label.jpg")})
    variants = [{"variant_id": f"C{i}", "name": f"Colour {i}", "size_value": None, "size_unit": None, "market": None,
                 "prices": [], "evidence": EV} for i in range(8)]
    rec, notes = normalise(ProductRecord.model_validate(_record(cautions=cautions, variants=variants)))
    uncovered = [{"kind": k, "text": t, "file": "spec.txt", "locator": "p2", "quote": q, "record_has": None,
                  "in_file_text": True} for k, t, q in (
        ("number", "size 2", "Size chart: S 60-64 cm"), ("number", "size 2 again", "Size chart: S 60-64 cm"),
        ("condition", "one per order", "Limit one per order"), ("caution", "held", "Keep away from eyes."))]
    result = OnboardingResult(record=rec, method="gemini_whole", seconds=1, notes=notes,
                              caution_pass_added=[c["text"] for c in cautions], uncovered=uncovered)
    risk = check_items(result, "risk", law="15 rules", caution_pass_only=result.caution_pass_added)
    checks, folded, auto = lean_items(result, risk, {"spec.txt": spec, "label.jpg": ""})
    assert len(checks) <= LEAN_MAX < len(risk)
    assert [a["label"] for a in auto] == [c["text"] for c in cautions[:4]]  # quotes found in the file text
    assert any(x["label"] == "Not for children." for x in checks + folded)  # an image quote stays to check
    cov = [x["label"] for x in checks + folded if x["group"] == "coverage"]
    assert cov == ["one per order", "size 2"]  # no duplicate, no line the record quotes, conditions first
    product = [x for x in checks if x["group"] == "product"]
    assert [x["label"] for x in product if "variants" in x["label"]] == ["8 variants: no price for this size"]
    assert [x["group"] for x in checks + folded].count("audit") <= 1


def test_d255_pack_totals_are_computed_in_code():
    from app.onboarding.release import release_pack

    ev = [{"file": "f", "locator": "l", "quote": "Box of 10 bottles; collagen 5,000 mg per bottle"}]
    rec, notes = normalise(ProductRecord.model_validate(_record(
        variants=[{"variant_id": "BOX", "name": "Collagen box", "size_value": None, "size_unit": None, "market": None,
                   "pack_count": 10, "pack_unit": "bottles",
                   "prices": [{"currency": "SGD", "minor": 3600, "kind": "live", "market": "SG", "evidence": EV}],
                   "evidence": ev}],
        ingredients=[{"name": "Collagen", "amount_value": "5000", "amount_unit": "mg", "per": "bottle",
                      "purpose": None, "evidence": ev}])))
    assert rec.variants[0].pack_count == 10
    ctx = make_ctx()
    try:
        facts = [p.text for p in release_pack(ctx.pack, rec, "rel_x_v1", None).passages if p.text.startswith("Computed")]
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    assert facts == ["Computed: Collagen box (10 × bottles): Collagen 5000 mg per bottle, 50000 mg in the whole pack"]


async def test_d255_an_absence_note_and_another_substance_are_no_corrections(tmp_path, monkeypatch):
    facts = [{"text": "Do not towel off after swimming or water exposure.", "evidence": EV},
             {"text": "Zinc 2.5 mg (as zinc gluconate) per bottle.", "evidence": EV}]

    def note(pid, ps):
        p = next(p for p in ps if p.text == facts[0]["text"])
        return _claim(pid, p, field="usage", stated_value="reapply after swimming",
                      correct_value="Not water-resistant, reapply hourly after swimming")

    row: dict = {}
    assert await _release_line(tmp_path / "a", monkeypatch, {"key_facts": facts},
                               "Reapply after swimming, and also every two hours.", note, row) == []
    assert [d["reason"] for d in row["dropped"]] == ["L1: an absence the quote does not state"]

    def salt(pid, ps):
        p = next(p for p in ps if p.text == facts[1]["text"])
        return _claim(pid, p, field="ingredient_amount", stated_value="3 mg zinc gluconate", correct_value="2.5 mg zinc")

    row = {}
    assert await _release_line(tmp_path / "b", monkeypatch, {"key_facts": facts},
                               "Each bottle has three mg zinc gluconate.", salt, row) == []
    assert [d["reason"] for d in row["dropped"]] == ["another substance than the host named"]


async def test_d256_an_audience_is_no_category_and_an_uncovered_type_maps_to_a_covered_one(tmp_path):
    from app.agents.inputs import current_product_json
    from app.logic.catalog import Catalog
    from app.onboarding.checks import check_items
    from app.onboarding.read import GENERIC_CATEGORIES, category_list, covered_categories
    from app.onboarding.release import release_pack

    assert "baby" not in GENERIC_CATEGORIES
    backend = ScriptedBackend()
    backend.outputs[("onboarding_read", "spec.txt")] = _record(category="baby", target_group=["infants"],
                                                               target_group_evidence=EV)
    backend.outputs[("onboarding_category", "category|baby")] = {"category": "household"}
    backend.outputs[("onboarding_category", "covered|household")] = {"category": "body_care"}
    ctx = make_ctx(backend=backend)
    try:
        cats, covered = category_list(ctx.pack, tmp_path), covered_categories(ctx.pack, tmp_path)
        res = await read_product(ctx.llm, [model_part("spec.txt", b"hello")], categories=cats, covered=covered)
        r = res.record
        assert r.category == "body_care" and r.target_group == ["infants"]
        assert r.category_evidence[-1].quote == "baby"
        pack = release_pack(ctx.pack, r, "rel_x_v1", None)
        assert len(pack.rules.rules) > 1  # the body-care law pack, not the qualifier rule alone
        catalog = Catalog(pack, "rel_x_v1")
        assert current_product_json(catalog, catalog.order[0])["target_group"] == ["infants"]
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    gov = [x["label"] for x in check_items(res, "risk", law="15 rules") if x["group"] == "governance"]
    assert "household has no law pack; mapped to the covered type body_care" in gov


def test_d257_test_conditions_reach_the_say_line_and_the_fact_line():
    from app.onboarding.release import release_pack

    result = [{"statement": "Stays on 8 hours", "numbers": ["8"], "provenance": "in_house", "qualifier": "in our test",
               "marker": None, "conditions": "indoors, on dry skin", "evidence": EV}]
    rec, _ = normalise(ProductRecord.model_validate(_record(qualified_results=result)))
    ctx = make_ctx()
    try:
        pack = release_pack(ctx.pack, rec, "rel_x_v1", None)
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    [q] = pack.products.products[0].qualified_results
    assert q.say == "Stays on 8 hours, indoors, on dry skin (in our test)"
    assert any(p.text.endswith("Holds only under: indoors, on dry skin.") for p in pack.passages)


def test_d257_lean_never_folds_what_the_qualifier_needs():
    from app.onboarding.checks import LEAN_MAX, lean_items
    from app.onboarding.record import OnboardingResult

    result = [{"statement": "Tear-free formula", "numbers": [], "provenance": "in_house", "qualifier": "in our eye test",
               "marker": None, "conditions": None, "evidence": EV}]
    rec, _ = normalise(ProductRecord.model_validate(_record(qualified_results=result)))
    res = OnboardingResult(record=rec, method="gemini_whole", seconds=1)
    risk = [{"group": "product", "label": f"p{i}", "reason": f"r{i}"} for i in range(LEAN_MAX + 3)]
    risk.append({"group": "safety", "label": "Keep the qualifier with the tear-free formula claim", "reason": "r"})
    shown, folded, _ = lean_items(res, risk, {})
    assert any("tear-free" in x["label"] for x in shown) and len(shown) == LEAN_MAX and len(folded) == 4


@pytest.mark.parametrize("line, alert", [
    ("Testing confirmed SPF50+ PA++++ protection for this sunscreen.", True),
    ("Both sizes deliver SPF50+ PA++++ according to the protection tests.", True),
    ("Larkspur Labs confirmed SPF50+ PA++++ protection.", False),
])
async def test_d258_a_capitalised_word_is_no_named_source(tmp_path, monkeypatch, line, alert):
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
    import eval_adversarial

    from app.onboarding.release import product_key, save_release

    from .judge_fakes import violation

    real_sleep = __import__("asyncio").sleep
    monkeypatch.setattr(__import__("asyncio"), "sleep", lambda _s: real_sleep(0))
    result = [{"statement": "SPF50+ PA++++", "numbers": ["50"], "provenance": "third_party_lab",
               "qualifier": "tested by Larkspur Labs in SPF and PA protection tests", "marker": None,
               "conditions": None, "evidence": EV}]
    rec, _ = normalise(ProductRecord.model_validate(_record(qualified_results=result)))
    rid = save_release(tmp_path, rec, "job")
    monkeypatch.setattr(eval_adversarial, "RELEASE", rid)
    backend = ScriptedBackend()
    backend.judge(line, violations=[violation("TEST-QUALIFIER", "SPF50+ PA++++", severity="MEDIUM")])
    ctx = make_ctx(backend=backend, ONBOARDING_DIR=str(tmp_path))
    case = {"case_id": "Q", "type": "t", "current_product_id": product_key(rec), "utterance": line,
            "expected": {"risky": True, "kind": "compliance"}}
    try:
        row = await eval_adversarial.run_case(ctx, case, 0)
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    assert bool([a for a in row["alerts"] if a["rule_id"] == "TEST-QUALIFIER"]) is alert


@pytest.mark.parametrize("line, stated, alert", [
    ("At least forty dollars qualifies.", "at least S$40", True),  # "above" is not "at least"
    ("Anything over forty dollars qualifies.", "over S$40", False),  # the same threshold: a restatement
])
async def test_d259_the_kind_of_threshold_is_compared(tmp_path, monkeypatch, line, stated, alert):
    facts = [{"text": "Free gift on orders above S$40.", "evidence": EV}]

    def claim(pid, ps):
        p = next(p for p in ps if p.text == facts[0]["text"])
        return _claim(pid, p, field="promo", stated_value=stated, correct_value="S$40")

    assert bool(await _release_line(tmp_path, monkeypatch, {"key_facts": facts}, line, claim)) is alert


def test_d259_offer_mixes_versions_and_unlisted_ingredients():
    from app.agents.fact_checker import _absent_from_sheet
    from app.models.data import Passage
    from app.onboarding.release import release_pack

    ev = [{"file": "f", "locator": "l", "quote": "3 for S$60"}]
    rec, _ = normalise(ProductRecord.model_validate(_record(
        offers=[{"offer_id": "o1", "kind": "fixed_price", "title": "3 for S$60", "variant_ids": [], "percent": None,
                 "amount": {"currency": "SGD", "minor": 6000, "kind": "promo", "market": "SG", "evidence": ev},
                 "buy": 3, "get": None, "gift": None, "rounding": None, "starts": None, "ends": None,
                 "duration_min": None, "evidence": ev}],
        conflicts=[{"field": "variants[0].size_value", "values": ["30", "35"], "evidence": EV,
                    "effective_dates": [None, None], "note": "the 30 ml sheet is archived"}])))
    ctx = make_ctx()
    try:
        pack = release_pack(ctx.pack, rec, "rel_x_v1", None)
        texts = [p.text for p in pack.passages]
        from app.agents.fact_checker import FactChecker
        from app.logic.catalog import Catalog

        fc = FactChecker(ctx.llm, ctx.prompts["fact_checker"], Catalog(pack, "rel_x_v1"))
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    [offer] = [t for t in texts if t.startswith("Offer: 3 for S$60")]
    assert "best price for 2 = S$72.00, 3 = S$60.00, 4 = S$96.00, 5 = S$132.00" in offer
    assert any(t.startswith("Versions of variants[0].size_value: 30 / 35 (the 30 ml sheet is archived)") for t in texts)
    assert "The show's language is en" in fc.system_instruction
    sheet = [Passage(passage_id="x", product_id="p", doc="d", section="Ingredients", text="Water, apple juice")]
    assert _absent_from_sheet("pear juice", sheet, "p") and not _absent_from_sheet("apple juice", sheet, "p")


def test_d259_a_negated_caution_is_corrected_by_that_caution(pack):
    from app.agents.utterance_judge import negated_caution
    from app.models.agent_io import FactCheckOutput

    p = next(x for x in pack.products.products if x.cautions)
    lex = pack.lexicon
    caution = p.cautions[0]
    assert negated_caution(f"No need to worry, {caution.lower()} is not something you must follow", p, lex) == caution
    assert negated_caution("It feels light and fresh on the skin.", p, lex) is None
    FactCheckOutput.model_validate({"claims": [{"span": "x", "field": "caution", "product_id": p.product_id,
                                                "stated_value": "no need to stop", "verdict": "contradicted",
                                                "correct_value": caution, "passage_id": None, "quote": None}],
                                    "reason": "t"})


async def test_d260_a_result_said_in_number_words_keeps_its_qualifier_alert(tmp_path, monkeypatch):
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
    import eval_adversarial

    from app.onboarding.release import product_key, save_release

    from .judge_fakes import violation

    real_sleep = __import__("asyncio").sleep
    monkeypatch.setattr(__import__("asyncio"), "sleep", lambda _s: real_sleep(0))
    line = "This tint stays non-transferring for eight hours."
    result = [{"statement": "Non-transferring for 8 hours", "numbers": ["8"], "provenance": "third_party_lab",
               "qualifier": "tested by Larkspur Labs, report LL-77", "marker": None, "conditions": None,
               "evidence": EV}]
    rec, _ = normalise(ProductRecord.model_validate(_record(qualified_results=result)))
    rid = save_release(tmp_path, rec, "job")
    monkeypatch.setattr(eval_adversarial, "RELEASE", rid)
    backend = ScriptedBackend()
    backend.judge(line, violations=[violation("TEST-QUALIFIER", "stays non-transferring for eight hours",
                                              severity="MEDIUM")])
    ctx = make_ctx(backend=backend, ONBOARDING_DIR=str(tmp_path))
    case = {"case_id": "Q", "type": "t", "current_product_id": product_key(rec), "utterance": line,
            "expected": {"risky": True, "kind": "compliance"}}
    try:
        row = await eval_adversarial.run_case(ctx, case, 0)
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()
    assert [a["rule_id"] for a in row["alerts"]] == ["TEST-QUALIFIER"] and row["dropped"] == []


async def test_d260_opposite_polarity_is_no_restatement_and_availability_is_a_field(tmp_path, monkeypatch):
    from app.models.agent_io import FactCheckOutput

    facts = [{"text": "Exactly S$40 does not qualify; orders above S$40 get the gift.", "evidence": EV}]

    def claim(pid, ps):
        p = next(p for p in ps if p.text == facts[0]["text"])
        return _claim(pid, p, field="promo", stated_value="exactly S$40 qualifies",
                      correct_value="Exactly S$40 does not qualify.", quote="Exactly S$40 does not qualify")

    assert await _release_line(tmp_path, monkeypatch, {"key_facts": facts}, "Exactly forty dollars qualifies.", claim)
    FactCheckOutput.model_validate({"claims": [{"span": "x", "field": "availability", "product_id": "p",
                                                "stated_value": "in stores", "verdict": "contradicted",
                                                "correct_value": "online only", "passage_id": None, "quote": None}],
                                    "reason": "t"})


async def test_d260_a_failed_compliance_call_is_made_once_more():
    from app.pipelines.utterances import UtterancePipeline
    from app.providers.llm import AgentResult

    calls = []

    async def call():
        calls.append(1)
        return AgentResult("timeout" if len(calls) == 1 else "ok", None, 10, "c", None, {})

    res = await UtterancePipeline._once_more(call)
    assert res.status == "ok" and len(calls) == 2


@pytest.mark.parametrize("numbers", [["8"], [], ["8.0"]])
def test_d261_a_result_matches_by_its_statement_numbers_too(pack, numbers):
    from app.logic import qualifier
    from app.models.data import QualifiedResult

    r = QualifiedResult(result_id="x", numbers=numbers,
                        say="Non-transferring for 8 hours, indoors (tested by Larkspur Labs, report LL-77)")
    p = pack.products.products[0].model_copy(update={"qualified_results": [r]})
    assert qualifier.approved_replacement("this tint stays non-transferring for eight hours", p, [], pack.lexicon) == r


async def test_d261_a_word_correction_the_line_already_says_is_dropped(tmp_path, monkeypatch):
    facts = [{"text": "Available online only.", "evidence": EV},
             {"text": "Free gift on orders above S$40.", "evidence": EV}]

    def online(pid, ps):
        p = next(p for p in ps if p.text == facts[0]["text"])
        return _claim(pid, p, field="availability", stated_value="available online only in both markets",
                      correct_value="Online only")

    row: dict = {}
    assert await _release_line(tmp_path / "a", monkeypatch, {"key_facts": facts},
                               "It is available online only, in both listed markets.", online, row) == []
    assert [d["reason"] for d in row["dropped"]] == ["D-261: the line already says the correction"]

    def threshold(pid, ps):
        p = next(p for p in ps if p.text == facts[1]["text"])
        return _claim(pid, p, field="promo", stated_value="S$40 does not qualify (threshold above S$40)",
                      correct_value="above S$40")

    row = {}
    assert await _release_line(tmp_path / "b", monkeypatch, {"key_facts": facts},
                               "Forty-four dollars qualifies while forty does not.", threshold, row) == []


def test_d262_a_correction_without_the_name_and_multiplicative_words(pack):
    from app.agents.fact_checker import _same_substance

    assert _same_substance("1% shea butter", "10% by weight", "shea butter: 10% by weight")
    assert _same_substance("9 mg sodium", "90 mg per 30 g scoop", "sodium: 90 mg per 30 g scoop")
    assert not _same_substance("3 mg zinc gluconate", "2.5 mg zinc", "Zinc 2.5 mg (as zinc gluconate)")
    assert pack.lexicon.times_numbers("Brush twice daily.") == {"2"}


async def test_d262_twice_is_a_number_in_every_guard(tmp_path, monkeypatch):
    facts = [{"text": "Brush twice daily.", "evidence": EV}]

    def claim(pid, ps):
        p = next(p for p in ps if p.text == facts[0]["text"])
        return _claim(pid, p, field="usage", stated_value="three times daily", correct_value="Brush twice daily.")

    [a] = await _release_line(tmp_path, monkeypatch, {"how_to_use": facts}, "Brush three times daily.", claim)
    assert a["correct_value"] == "Brush twice daily."


def test_d263_the_bucket_holds_onboarding_state_across_a_restart(tmp_path):
    from app.onboarding import storage
    from app.onboarding.files import store
    from app.onboarding.lawpack import LawPackRecord, lawpack_for, save_lawpack
    from app.onboarding.release import load_release_full, save_release
    from app.onboarding.sources import release_sources

    bucket = storage.MemoryBucket()
    a, b = tmp_path / "a", tmp_path / "b"
    storage.configure(a, bucket)
    try:
        rec, _ = normalise(ProductRecord.model_validate(_record()))
        stored = store(a, "job1", "spec.txt", b"Use within 12 months.")
        rid = save_release(a, rec, "job1")
        law = LawPackRecord.model_validate({"name": "Toy law", "regulator": "R", "covers": ["toy"], "language": "en",
                                           "rules": []})
        pid = save_lawpack(a, "SG", law, "job2")
        keys = set(bucket.keys(storage.PREFIX))
        assert {f"onboarding/releases/{rid}.json", f"onboarding/lawpacks/{pid}.json", "onboarding/job1/files.json",
                f"onboarding/job1/{stored.path.name}"} <= keys  # write-through
        # write-once: another release under the same name is refused
        (a / "releases" / f"{rid}.json").write_text("{}", encoding="utf-8")
        with pytest.raises(storage.Immutable):
            storage.persisted(a / "releases" / f"{rid}.json", create_only=True)

        storage.configure(b, bucket)  # a cold start: an empty disk, the same bucket
        assert load_release_full(b, rid)[0].product_name == rec.product_name  # read back at start
        assert lawpack_for(b, "SG", "toy")[0] == pid  # listings work from the read-back state
        assert not (b / "job1" / stored.path.name).exists()  # uploads stay remote until needed
        assert release_sources(str(b), rid)[1][0].text.startswith("Use within 12 months.")  # read back on a miss
        assert save_release(b, rec, "job3") == rid.replace("_v1", "_v2")  # versions continue, never overwrite
    finally:
        storage.configure(a, None)
        storage.configure(b, None)


def test_d263_onboard_approve_restart_and_start_a_show_from_the_release(tmp_path, monkeypatch):
    import time

    from fastapi.testclient import TestClient

    from app.main import create_app
    from app.onboarding import storage

    from .helpers import make_settings

    bucket = storage.MemoryBucket()
    monkeypatch.setattr(storage, "bucket_for", lambda name, project=None: bucket)
    backend = ScriptedBackend()
    backend.outputs[("onboarding_read", "spec.txt")] = _record(brand="Lumen", product_name="Night Serum")
    first, second = tmp_path / "first", tmp_path / "second"
    try:
        with TestClient(create_app(make_settings(ONBOARDING_DIR=str(first), ONBOARDING_GCS_BUCKET="fake"),
                                   llm_backend=backend)) as c:
            job = c.post("/api/onboarding/jobs", files=[("files", ("spec.txt", b"hello", "text/plain"))]).json()["job_id"]
            for _ in range(100):
                if c.get(f"/api/onboarding/jobs/{job}").json()["status"] == "done":
                    break
                time.sleep(0.02)
            rid = c.post(f"/api/onboarding/jobs/{job}/approve", json={}).json()["release_id"]
        # a new instance after scale-to-zero: another (empty) disk, the same bucket
        with TestClient(create_app(make_settings(ONBOARDING_DIR=str(second), ONBOARDING_GCS_BUCKET="fake"),
                                   llm_backend=backend)) as c:
            listed = c.get("/api/onboarding/releases").json()  # D-264: the cold instance lists it
            assert [(x["release_id"], x["product_name"], x["brand"], x["version"]) for x in listed] == [
                (rid, "Night Serum", "Lumen", 1)]
            r = c.post("/api/sessions", json={"release_id": listed[0]["release_id"]})
            assert r.status_code == 201
            cat = c.get(f"/api/catalog?session={r.json()['session_id']}").json()
            assert [p["name"] for p in cat["products"]] == ["Night Serum"]
    finally:
        storage.configure(first, None)
        storage.configure(second, None)


@pytest.mark.parametrize("brief_ok", [True, False])
def test_d265_approval_briefs_the_release_for_stt_and_the_prompter(tmp_path, brief_ok):
    import time

    from fastapi.testclient import TestClient

    from app.api.audio_ws import release_phrases
    from app.logic.catalog import Catalog
    from app.main import create_app
    from app.onboarding.release import product_key

    from .helpers import make_settings

    record = _record(brand="Quillon", product_name="Kelpa Night Serum",
                     approved_claims=[{"text": "Helps skin look smoother by morning.", "evidence": EV}],
                     cautions=[{"text": "Patch test first.", "evidence": EV}])
    rid = f"rel_{product_key(normalise(ProductRecord.model_validate(record))[0])}_v1"
    backend = ScriptedBackend()
    backend.outputs[("onboarding_read", "spec.txt")] = record
    if brief_ok:
        backend.outputs[("brief_agent", rid)] = {"product_id": rid, "reason": "t", "glossary": [
            {"term": "Quillon", "kind": "brand", "spoken_forms": ["kwil-on", "quill on"], "source_section": "x"},
            {"term": "Kelpa", "kind": "product", "spoken_forms": ["kelp a"], "source_section": "x"},
            {"term": "Invented", "kind": "technical", "spoken_forms": [], "source_section": "x"}],  # not in the doc
            "prompter_points": [
                {"text": "Patch test first.", "required": True, "source_section": "Cautions"},
                {"text": "Helps skin look smoother by morning.", "required": False, "source_section": "Claims"}]}
    else:
        backend.fail("brief_agent", rid)
    with TestClient(create_app(make_settings(ONBOARDING_DIR=str(tmp_path)), llm_backend=backend)) as c:
        job = c.post("/api/onboarding/jobs", files=[("files", ("spec.txt", b"hello", "text/plain"))]).json()["job_id"]
        for _ in range(100):
            if c.get(f"/api/onboarding/jobs/{job}").json()["status"] == "done":
                break
            time.sleep(0.02)
        approved = c.post(f"/api/onboarding/jobs/{job}/approve", json={}).json()
        assert approved["release_id"] == rid and approved["live"]["brief"] == {"status": "running"}  # D-266
        for _ in range(100):
            status = c.get(f"/api/onboarding/releases/{rid}/brief").json()["status"]
            if status != "running":
                break
            time.sleep(0.02)
        assert status == ("done" if brief_ok else "failed (error)")
        assert c.get("/api/onboarding/releases").json()[0]["brief"] == status
        sid = c.post("/api/sessions", json={"release_id": rid}).json()["session_id"]
        state = c.post(f"/api/sessions/{sid}/start") and c.get(f"/api/sessions/{sid}/state?role=host").json()
        ctx = c.app.state.ctx
        phrases = release_phrases(ctx, Catalog(ctx.pack_for("SG", ctx.pack.lang, rid)[0], rid), rid)
    points = [(p["text"], p["required"]) for p in state["prompter"]["points"]]
    if brief_ok:
        assert (tmp_path / "briefs" / f"{rid}.json").is_file()
        assert points == [("Patch test first.", True), ("Helps skin look smoother by morning.", False)]
        assert phrases[:4] == ["Quillon", "kwil-on", "quill on", "Kelpa"] and "Invented" not in phrases
    else:
        assert not (tmp_path / "briefs").exists()
        assert points == [("Helps skin look smoother by morning.", True)]  # today's behaviour
        assert phrases[0] == "Quillon"


def test_d266_a_show_started_before_the_brief_lands_keeps_todays_points_and_the_next_gets_the_brief(tmp_path):
    import time

    from fastapi.testclient import TestClient

    from app.main import create_app
    from app.onboarding.release import product_key

    from .helpers import make_settings

    record = _record(approved_claims=[{"text": "Helps skin look smoother by morning.", "evidence": EV}])
    rid = f"rel_{product_key(normalise(ProductRecord.model_validate(record))[0])}_v1"
    backend = ScriptedBackend()
    backend.outputs[("onboarding_read", "spec.txt")] = record
    backend.outputs[("brief_agent", rid)] = {"product_id": rid, "reason": "t", "glossary": [], "prompter_points": [
        {"text": "Patch test first.", "required": True, "source_section": "Cautions"}]}
    backend.delays[("brief_agent", rid)] = 0.5

    def points(c, sid):
        c.post(f"/api/sessions/{sid}/start")
        return [p["text"] for p in c.get(f"/api/sessions/{sid}/state?role=host").json()["prompter"]["points"]]

    with TestClient(create_app(make_settings(ONBOARDING_DIR=str(tmp_path)), llm_backend=backend)) as c:
        job = c.post("/api/onboarding/jobs", files=[("files", ("spec.txt", b"hello", "text/plain"))]).json()["job_id"]
        for _ in range(100):
            if c.get(f"/api/onboarding/jobs/{job}").json()["status"] == "done":
                break
            time.sleep(0.02)
        t0 = time.perf_counter()
        c.post(f"/api/onboarding/jobs/{job}/approve", json={})
        assert time.perf_counter() - t0 < 0.5  # the approval does not wait for the brief
        early = c.post("/api/sessions", json={"release_id": rid}).json()["session_id"]
        assert c.get(f"/api/onboarding/releases/{rid}/brief").json()["status"] == "running"
        assert points(c, early) == ["Helps skin look smoother by morning."]
        for _ in range(100):
            if c.get(f"/api/onboarding/releases/{rid}/brief").json()["status"] == "done":
                break
            time.sleep(0.02)
        late = c.post("/api/sessions", json={"release_id": rid}).json()["session_id"]
        assert points(c, late) == ["Patch test first."]


def test_d269_brief_offers_terms_arm_gives_every_offer_term_and_room_for_it():
    from app.agents.brief import post_validate
    from types import SimpleNamespace

    from app.models.agent_io import BriefOutput
    from app.onboarding.brief import release_document

    ev = [{"file": "f", "locator": "l", "quote": "Code SAVE10: 10% off, online only, one per order"}]
    rec = ProductRecord.model_validate(_record(offers=[
        {"offer_id": "o1", "kind": "percent_off", "title": "Code SAVE10", "variant_ids": [], "percent": "10",
         "amount": None, "buy": None, "get": None, "gift": None, "rounding": None, "starts": "2026-11-01",
         "ends": "2026-11-14", "duration_min": None, "combinable": False, "conditions": ["one per order"],
         "evidence": ev}]))
    sheet, terms = release_document(rec), release_document(rec, offers="terms")
    assert "2026-11-14" not in sheet and "combinable" not in sheet  # today's sheet drops them
    line = next(x for x in terms.splitlines() if x.startswith("- [o1]"))
    for want in ("10% off", "2026-11-01", "2026-11-14", "combinable with other offers or codes: no",
                 "condition: one per order", "online only"):
        assert want in line
    long = "Code SAVE10: 10% off, online store only, 1 to 14 Nov, not with other offers, one per order. " * 2
    out = BriefOutput(product_id="r", glossary=[], reason="t", prompter_points=[
        {"text": long, "required": True, "source_section": "Offer o1"},
        {"text": long, "required": False, "source_section": "Claims"}])
    reflex = SimpleNamespace(scan=lambda text: [])  # no restricted wording here
    assert [p.source_section for p in post_validate(out, "r", "", reflex, 10, 140).prompter_points] == []
    assert [p.source_section for p in post_validate(out, "r", "", reflex, 10, 140, 240).prompter_points] == ["Offer o1"]


@pytest.mark.parametrize("wording", ["d256", "split"])
def test_d270_vulnerable_wording_arm_swaps_only_the_target_group_bullet(wording):
    ctx = make_ctx(VULNERABLE_WORDING=wording)
    try:
        for name in ("compliance_judge_rules", "utterance_judge", "judge_single"):  # D-358: the rules prompt
            p = ctx.prompts[name]
            assert ("suitability promise" in p.system_instruction) is (wording == "d256"), name
            assert ("which ages may use it" in p.system_instruction) is (wording == "split"), name
            assert p.version.endswith("+vs") is (wording == "split")
            assert "- `is_quoted = true`" in p.system_instruction  # the next bullet is kept
            assert "skin types" in p.system_instruction  # D-256's skin-type bullet is kept
    finally:
        ctx.runtimes.shutdown()
        ctx.scheduler.shutdown()


def test_d271_comparison_kind_takes_the_longest_phrase(pack):
    lex = pack.lexicon
    assert lex.comparison_kind("S$40 and above") == "ge"  # "above" alone is gt
    assert lex.comparison_kind("for children 4 and up") == "ge"
    assert lex.comparison_kind("ages 12 and under") == "le"
    assert lex.comparison_kind("above S$40") == "gt"


@pytest.mark.parametrize("line, stated, dropped", [
    ("For children 4 and up.", "children 4 and up", True),  # the same range as "not ... under 4"
    ("For children 3 and up.", "children 3 and up", False),  # another number
    ("For children under 4 too.", "children under 4", False),  # the opposite range
])
async def test_d271_an_age_band_that_states_the_cautions_range_is_no_contradiction(tmp_path, monkeypatch, line,
                                                                                   stated, dropped):
    facts = [{"text": "Not suitable for children under 4 years.", "evidence": EV}]

    def claim(pid, ps):
        p = next(p for p in ps if p.text == facts[0]["text"])
        return _claim(pid, p, field="caution", stated_value=stated, correct_value=facts[0]["text"])

    row: dict = {}
    alerts = await _release_line(tmp_path, monkeypatch, {"cautions": facts}, line, claim, row)
    assert (alerts == []) is dropped


async def test_d272_offer_points_come_from_the_record_and_survive_translation_only_intact():
    from types import SimpleNamespace

    from app.onboarding.brief import offer_lines, offer_points

    money = {"currency": "SGD", "minor": 3000, "kind": "promo", "market": "SG", "evidence": EV}
    rec = ProductRecord.model_validate(_record(offers=[
        {"offer_id": "o1", "kind": "percent_off", "title": "Code SAVE10", "variant_ids": [], "percent": "10",
         "amount": None, "buy": None, "get": None, "gift": None, "rounding": None, "starts": "2026-11-01",
         "ends": "2026-11-14", "duration_min": None, "combinable": False,
         "conditions": ["online store only", "one per order"], "evidence": EV},
        {"offer_id": "o2", "kind": "spend_threshold_gift", "title": "Free pouch", "variant_ids": [], "percent": None,
         "amount": money, "buy": None, "get": None, "gift": "pouch", "rounding": None, "starts": None, "ends": None,
         "duration_min": None, "combinable": None, "conditions": ["members only"], "evidence": EV}]))
    assert offer_lines(rec) == [
        ("o1", ["Code SAVE10: 10% off, from 1 to 14 November 2026.", "Terms: online store only; one per order."]),
        ("o2", ["Free pouch: when you spend the S$30 spend in the terms.", "Terms: members only."])]
    class Fake:
        async def call(self, req):
            lines = json.loads(req.user_message.split('<data name="LINES">')[1].split("</data>")[0])["lines"]
            out = [x.replace("from", "dari").replace("S$30", "S$3") for x in lines]  # one number broken
            return SimpleNamespace(status="ok", output=SimpleNamespace(lines=out))

    en = [p.text for p in await offer_points(Fake(), rec, "en")]  # an English record and show: no call
    ms = [p.text for p in await offer_points(Fake(), rec, "ms")]
    assert en[0] == "Code SAVE10: 10% off, from 1 to 14 November 2026." and ms[0] == "Code SAVE10: 10% off, dari 1 to 14 November 2026."
    assert ms[2] == "Free pouch: when you spend the S$30 spend in the terms."  # a changed number keeps the English


async def test_d273_a_release_without_approved_claims_keeps_the_fallback_alert(tmp_path, monkeypatch):
    """A failed compliance call on a product with no approved claim raised IndexError in the say-instead fallback,
    so the background task died and the line's red alert was lost."""
    row: dict = {}
    await _release_line(tmp_path, monkeypatch, {}, "It cures eczema.", lambda pid, ps: _claim(
        pid, ps[0], field="other", verdict="not_in_sources", stated_value="x", correct_value=None), row,
                        fail=("compliance_judge", "utterance_judge"))
    [a] = [a for a in row["alerts"] if a["kind"] == "compliance"]
    assert a["severity"] == "HIGH" and a["say_instead"] == "(no approved wording in en)"


async def test_d274_cautions_and_directions_in_the_show_language_and_repeats_merged():
    from types import SimpleNamespace

    from app.onboarding.translate import show_language_lines

    ev = lambda q: [{"file": "f", "locator": "l", "quote": q}]  # noqa: E731
    rec = ProductRecord.model_validate(_record(
        cautions=[{"text": "Keep out of reach of children.", "evidence": ev("Keep out of reach of children.")},
                  {"text": "请放在儿童接触不到的地方。", "evidence": ev("请放在儿童接触不到的地方。")},
                  {"text": "勿超过每日2粒。", "evidence": ev("勿超过每日2粒。")}],
        how_to_use=[{"text": "每2小时补涂一次。", "evidence": ev("每2小时补涂一次。")}]))
    answer = {"cautions-0": ("Keep out of reach of children.", None),
              "cautions-1": ("Keep out of reach of children.", "cautions-0"),  # the same caution in Chinese
              "cautions-2": ("Do not exceed the daily dose.", None),  # the number was lost: kept as written
              "how_to_use-0": ("Reapply every 2 hours.", None)}

    class Fake:
        async def call(self, req):
            return SimpleNamespace(status="ok", output=SimpleNamespace(items=[
                SimpleNamespace(id=k, text=t, same_as=s) for k, (t, s) in answer.items()]))

    notes = await show_language_lines(Fake(), rec, "en")
    assert [c.text for c in rec.cautions] == ["Keep out of reach of children.", "勿超过每日2粒。"]
    assert [e.quote for e in rec.cautions[0].evidence] == ["Keep out of reach of children.", "请放在儿童接触不到的地方。"]
    assert [h.text for h in rec.how_to_use] == ["Reapply every 2 hours."]
    assert [h.evidence[0].quote for h in rec.how_to_use] == ["每2小时补涂一次。"]  # the original stays the evidence
    assert notes == ["show language: 1 cautions or directions put into en, 1 repeats merged",
                     "show language: 1 lines kept as written (a number changed in translation)"]


def test_d272_patch_offer_wording_follows_the_kind_and_the_terms_as_written(pack):
    from app.onboarding.brief import offer_lines

    ev = lambda q: [{"file": "f", "locator": "l", "quote": q}]  # noqa: E731
    money = {"currency": "SGD", "minor": 4000, "kind": "promo", "market": "SG", "evidence": EV}
    rec = ProductRecord.model_validate(_record(offers=[
        {"offer_id": "g", "kind": "spend_threshold_gift", "title": "Free sleeve", "variant_ids": [], "percent": None,
         "amount": money, "buy": 1, "get": 1, "gift": "travel sleeve", "rounding": None, "starts": "2026-10-01T00:00:00+08:00",
         "ends": "2026-10-31T23:59:00+08:00", "duration_min": None, "combinable": True,
         "conditions": ["combinable only with the member price"], "evidence": ev("orders above S$40 get a free sleeve")}]))
    [(_, lines)] = offer_lines(rec, pack.lexicon)
    assert lines == ["Free sleeve: get a free travel sleeve, when you spend above S$40, from 1 to 31 October 2026.",
                     "Terms: combinable only with the member price."]
    assert not any("buy 1" in x or "combinable with other" in x or "00:00" in x for x in lines)


def test_d272_readability_no_repeats_a_subject_document_notes_dropped_two_terms_lines():
    from app.onboarding.brief import offer_lines

    money = {"currency": "SGD", "minor": 400, "kind": "promo", "market": "SG", "evidence": EV}
    note = "In this bulletin, S0 means SGD 0.00"
    long = [f"limit number {i} applies to every order placed in the store during the offer week" for i in range(6)]
    rec = ProductRecord.model_validate(_record(offers=[
        {"offer_id": "a", "kind": "amount_off", "title": "S$4.00 off the second tube", "variant_ids": [],
         "percent": None, "amount": money, "buy": None, "get": None, "gift": None, "rounding": None, "starts": None,
         "ends": None, "duration_min": None, "conditions": ["is not combinable with codes", note, *long],
         "evidence": EV}]))
    [(_, lines)] = offer_lines(rec, None, notes=[note])
    assert lines[0] == "S$4.00 off the second tube."  # the amount the title says is not said again
    assert lines[1].startswith("Terms: It is not combinable with codes;")
    assert len(lines) == 3 and not any("bulletin" in x for x in lines)


def test_d276_every_record_caution_gets_a_required_point_worded_only_if_its_meaning_survives(pack):
    from app.models.agent_io import BriefOutput
    from app.onboarding.brief import _with_cautions

    t = lambda x: {"text": x, "evidence": EV}  # noqa: E731
    rec = ProductRecord.model_validate(_record(
        cautions=[t("Keep out of reach of children under 6."), t("Do not use on broken skin."),
                  t("Sold online only in Singapore."), t("Contains milk.")],
        how_to_use=[t("Shake well before use."), t("Store below 30 degrees."), t("Contains milk.")]))
    out = BriefOutput(product_id="r", glossary=[], reason="t", prompter_points=[
        {"text": "Softens dry lips.", "required": False, "source_section": "Approved on-air claims"},
        {"text": "Careful with kids!", "required": True, "source_section": "Cautions"}], cautions=[
        {"id": "c0", "kind": "safety", "say": "Keep it out of reach of children under 6."},
        {"id": "c1", "kind": "safety", "say": "Use it on broken skin."},  # the negation is lost
        {"id": "c2", "kind": "limit", "say": "Online only in Singapore."},
        {"id": "d0", "kind": "direction", "say": "Shake it."},
        {"id": "d1", "kind": "safety", "say": "Store it below 30 degrees."},
        {"id": "d2", "kind": "safety", "say": "Contains milk."}])  # c3 left out by the model; d2 repeats it
    points = _with_cautions(out, rec, "en", pack.lexicon)
    assert [(p.text, p.required) for p in points] == [
        ("Softens dry lips.", False),
        ("Keep it out of reach of children under 6.", True),
        ("Do not use on broken skin.", True),  # the record's text
        ("Contains milk.", True),  # never lost, even when the model leaves it out
        ("Store it below 30 degrees.", True)]

"""AI Builder Cup 2026 submission deck: the organizer's template filled with the current service (10/11).
Template: "Submission Template _ AI Builder Cup.pptx" at the repo root (the organizer's file, not in git).
Layout helpers come from the 9/26 proposal builder (pitch/build_submission_pptx.py in the planning repo);
content is the current service; screenshots are from the deployed app (docs/decks/img/sub1011).
Run: uv run --no-project --with python-pptx --with pillow python docs/decks/build/build_submission_1011.py"""
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[3]
SRC = ROOT / "Submission Template _ AI Builder Cup.pptx"
IMG = ROOT / "docs" / "decks" / "img" / "sub1011"
OUT = ROOT / "docs" / "decks" / "submission_1011.pptx"

INK, SUB, LINE, BUBBLE = "202124", "5F6368", "BDC1C6", "F1F3F4"
BLUE, RED, YELT, GRN = "1A73E8", "D93025", "B06000", "1E8E3E"
BLUEBG, REDBG, YELBG, GRNBG = "E8F0FE", "FCE8E6", "FEF7E0", "E6F4EA"
FONT = "Google Sans Flex"  # the template's font (OFL); render needs static faces cut from the google/fonts variable file
TONE = {"blue": (BLUE, BLUEBG), "red": (RED, REDBG), "yellow": (YELT, YELBG), "green": (GRN, GRNBG), "grey": (SUB, BUBBLE)}
SUP = {"sup": True, "color": BLUE, "bold": True}
FN = {"color": SUB}


# ---------------- icons (drawn once with PIL) ----------------
def _icon(kind):
    n = 1024
    im = Image.new("RGBA", (n, n), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    bg = {"person": "#5F6368", "ai": "#E8F0FE", "check": "#1E8E3E", "x": "#D93025", "plus": "#1A73E8",
          "warn": "#B06000"}[kind]
    d.ellipse([0, 0, n - 1, n - 1], fill=bg)
    if kind == "person":
        d.ellipse([352, 200, 672, 520], fill="white")
        d.pieslice([232, 580, 792, 1100], 180, 360, fill="white")
    elif kind == "ai":
        def star(cx, cy, r, w):
            d.polygon([(cx, cy - r), (cx + w, cy - w), (cx + r, cy), (cx + w, cy + w), (cx, cy + r), (cx - w, cy + w),
                       (cx - r, cy), (cx - w, cy - w)], fill="#1A73E8")
        star(450, 560, 300, 70)
        star(740, 300, 150, 36)
    elif kind == "check":
        d.line([(270, 530), (440, 700), (770, 360)], fill="white", width=110, joint="curve")
    elif kind == "plus":
        d.line([(512, 280), (512, 744)], fill="white", width=110)
        d.line([(280, 512), (744, 512)], fill="white", width=110)
    elif kind == "x":
        d.line([(320, 320), (704, 704)], fill="white", width=110)
        d.line([(704, 320), (320, 704)], fill="white", width=110)
    elif kind == "warn":
        d.line([(512, 250), (512, 600)], fill="white", width=120)
        d.ellipse([452, 690, 572, 810], fill="white")
    buf = BytesIO()
    im.resize((256, 256), Image.LANCZOS).save(buf, "PNG")
    return buf.getvalue()


ICONS = {k: _icon(k) for k in ("person", "ai", "check", "x", "plus", "warn")}


# ---------------- drawing helpers (inches) ----------------
def rgb(h):
    return RGBColor.from_string(h)


def _runs(p, spec, size, color, bold):
    """spec: str or list of (text, {size,color,bold,sup,strike,link}) runs."""
    for item in ([spec] if isinstance(spec, str) else spec):
        text, o = (item, {}) if isinstance(item, str) else item
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(o.get("size", size))
        f.name, f.bold = FONT, o.get("bold", bold)
        f.color.rgb = rgb(o.get("color", color))
        if o.get("sup"):
            r._r.get_or_add_rPr().set("baseline", "30000")
        if o.get("strike"):
            r._r.get_or_add_rPr().set("strike", "sngStrike")
        if o.get("link"):
            r.hyperlink.address = o["link"]
            f.color.rgb = rgb(BLUE)


def fill_text(tf, paras, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, space=0):
    tf.word_wrap = True
    tf.clear()
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space and i:
            p.space_before = Pt(space)
        _runs(p, spec, size, color, bold)


def tb(s, x, y, w, h, paras, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = t.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = anchor
    fill_text(tf, paras, size, color, bold, align, space)
    return t


def shape(s, kind, x, y, w, h, fill=None, border=None, radius=0.12, lw=0.75):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if border:
        sh.line.color.rgb = rgb(border)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def box(s, x, y, w, h, paras, fill=None, border=None, size=12, color=INK, bold=False, pad=0.1, anchor=MSO_ANCHOR.TOP,
        align=PP_ALIGN.LEFT, radius=0.12, space=0, lw=0.75):
    sh = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, border, radius, lw)
    tf = sh.text_frame
    for side in ("left", "right"):
        setattr(tf, f"margin_{side}", Inches(pad))
    for side in ("top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(min(pad, 0.06)))
    tf.vertical_anchor = anchor
    fill_text(tf, paras, size, color, bold, align, space)
    return sh


def icon(s, kind, x, y, d=0.4):
    s.shapes.add_picture(BytesIO(ICONS[kind]), Inches(x), Inches(y), Inches(d), Inches(d))


def arrow(s, x, y, w=0.3, color="9AA0A6"):
    return shape(s, MSO_SHAPE.RIGHT_ARROW, x, y - 0.06, w, 0.12, fill=color)


def down_arrow(s, x, y, h=0.22):
    return shape(s, MSO_SHAPE.DOWN_ARROW, x - 0.06, y, 0.12, h, fill="9AA0A6")


def hline(s, x1, y, x2, color=LINE, width=0.75):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    spPr = c._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    return c


MEASURE = ImageFont.truetype("arial.ttf", 100)  # Google Sans is about 8% wider (chip_w)


def chip_w(text, size):
    return MEASURE.getlength(text) * size / 100 / 72 * 1.08 + 0.26


def chip(s, x, y, text, tone="blue", size=9.5, w=None):
    fg, bg = TONE[tone]
    w = w or chip_w(text, size)
    sh = box(s, x, y, w, size / 72 * 1.9, [text], fill=bg, size=size, color=fg, pad=0.1, anchor=MSO_ANCHOR.MIDDLE,
             align=PP_ALIGN.CENTER, radius=0.5)
    sh.text_frame.word_wrap = False
    return w


def headline(s, text):
    tb(s, 0.4, 1.12, 9.2, 0.42, [text], size=18, bold=True)


def label(s, y, text, x=0.4, w=9.2, color=BLUE):
    tb(s, x, y, w, 0.24, [[(text, {"bold": True})]], size=10.5, color=color)


def cols(n, x0=0.4, width=9.2, gap=0.2):
    w = (width - (n - 1) * gap) / n
    return [(x0 + i * (w + gap), w) for i in range(n)]


def num_circle(s, x, y, n, d=0.28, color=BLUE):
    c = shape(s, MSO_SHAPE.OVAL, x, y, d, d, fill=color)
    tf = c.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_text(tf, [[(str(n), {"bold": True})]], size=9.5, color="FFFFFF", align=PP_ALIGN.CENTER)


def table(s, y, colspec, head, rows, row_h, size=10, head_size=9):
    """colspec: list of (x, w). A cell is a str, a run list, or ('chip', text, tone)."""
    for (x, w), hd in zip(colspec, head):
        tb(s, x, y, w, 0.25, [hd], size=head_size, color=SUB)
    x0, x1 = colspec[0][0], colspec[-1][0] + colspec[-1][1]
    hline(s, x0, y + 0.27, x1, width=1.25)
    yy = y + 0.33
    for r in rows:
        for (x, w), c in zip(colspec, r):
            if isinstance(c, tuple):
                chip(s, x, yy + 0.01, c[1], c[2], size=8.5)
            else:
                tb(s, x, yy, w - 0.1, row_h - 0.06, [c], size=size)
        yy += row_h
        hline(s, x0, yy - 0.04, x1)
    return yy


def picture(s, name, x, y, w=None, h=None, border=True):
    path = IMG / f"{name}.png"
    pw, ph = Image.open(path).size
    if w and not h:
        h = w * ph / pw
    elif h and not w:
        w = h * pw / ph
    s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if border:
        shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, border=LINE)
    return w, h


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def strip_body(s, title):
    """Keep only the template's title line in its text box (drop the prompt lines under it)."""
    tf = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith(title)).text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)


# ---------------- build ----------------
prs = Presentation(SRC)


def delete_slide(i):
    lst = prs.slides._sldIdLst
    sid = lst[i]
    prs.part.drop_rel(sid.rId)
    lst.remove(sid)


delete_slide(14)  # empty spare slide
delete_slide(1)   # "please download this template" note
S = list(prs.slides)  # 0 team, 1 brief, 2 opps, 3 features, 4 flow, 5 wire, 6 arch, 7 tech, 8 cost, 9 snaps, 10 bench, 11 more, 12 links, 13 thank you

# ---------------- 1 Team details ----------------
s = S[0]
tf = next(sh for sh in s.shapes if sh.has_text_frame and "Team Details" in sh.text_frame.text).text_frame
vals = {"Team name": "Uh?", "Team leader name": "Jaewoo Choi",
        "Problem Statement": "Live-commerce hosts speak unscripted for hours. A wrong price or a claim the regulations "
                             "forbid (\"this cream treats eczema\") can reach viewers before anyone corrects it."}
for p in tf.paragraphs:
    for k, v in vals.items():
        if p.text.strip().startswith(k):
            _runs(p, [(" " + v, {"bold": k != "Problem Statement"})], 13 if k == "Problem Statement" else 15, INK, False)
notes(s, "팀 정보. 문제 정의는 예전 제출본처럼 한 문장으로 쓰되, 지금 서비스(규정 위반 발언, 틀린 상품 정보)에 맞췄다.")

# ---------------- 2 Brief about the idea ----------------
s = S[1]
headline(s, "It checks what the host says, while they say it, and clears when they fix it")
tb(s, 0.4, 1.52, 9.2, 0.24, [[("A show on the deployed app (10/11), lines typed into the live judging path. Each final line goes to four Gemini judges at once.",
                              {})]], size=9, color=SUB)


def flow_row(y, said, judged, result, tone, kind, h=0.64):
    """[Host says] -> [Gemini judges] -> [screen shows], on one row."""
    icon(s, "person", 0.42, y + (h - 0.34) / 2, 0.34)
    box(s, 0.86, y, 3.0, h, [said], fill=BUBBLE, size=10.5, anchor=MSO_ANCHOR.MIDDLE, pad=0.12)
    arrow(s, 3.93, y + h / 2, 0.24)
    icon(s, "ai", 4.24, y + (h - 0.34) / 2, 0.34)
    box(s, 4.66, y, 2.3, h, judged, fill=BLUEBG, size=9.5, anchor=MSO_ANCHOR.MIDDLE, pad=0.12)
    arrow(s, 7.03, y + h / 2, 0.24)
    fg, bg = TONE[tone]
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.34, y, 2.26, h, fill=bg, border=fg if tone in ("red", "yellow") else None)
    icon(s, kind, 7.44, y + h / 2 - 0.13, 0.26)
    tb(s, 7.78, y, 1.78, h, result, size=9.5, color=fg, anchor=MSO_ANCHOR.MIDDLE)


rows2 = [
    ("\"And honestly, this cream treats eczema, so your eczema will be gone.\"",
     [[("Compliance judge", {"bold": True})], ["Disease claim: ASEAN cosmetic claims guideline, App. III 2.d"]],
     [[("Red box in 1.3 s", {"bold": True})], [("\"Sorry, I take that back. Helps relieve the feeling of dryness.\"", {"size": 8.5})]],
     "red", "x"),
    ("\"Today's live price is just S$32.\"",
     [[("Fact checker", {"bold": True})], ["The prompter shows S$36.00"]],
     [[("Orange box in 0.8 s", {"bold": True})], ["The live price is S$36.00."]], "yellow", "warn"),
    ("\"Sorry, I take that back. It helps relieve the feeling of dryness.\" · \"My mistake, the live price is S$36.\"",
     [[("Same judges, next lines", {"bold": True})], ["The claim is withdrawn; the right price is said"]],
     [[("Both boxes clear", {"bold": True})], ["in the host's own words"]], "green", "check"),
    ("\"It keeps your skin hydrated for up to 24 hours, in our in-house test.\" · \"Okay, next, let me show you the ampoule.\"",
     [[("Prompter and context judges", {"bold": True})], ["Full point said · \"next\" + the next product's category"]],
     [[("Point ticked · prompter moves on", {"bold": True})], ["no click by the host"]], "green", "check"),
]
for k, (said, judged, result, tone, kind) in enumerate(rows2):
    flow_row(1.84 + k * 0.73, said, judged, result, tone, kind)
box(s, 0.4, 4.8, 9.2, 0.52, [[("Viewer questions: ", {"bold": True}),
                               ("answered automatically when the product record answers them (\"how much is it today?\" got the live price in 4.3 s); "
                                "health questions (\"I get redness easily...\") go to the host's screen with a suggested reply.", {})]],
    fill="FFFFFF", border=LINE, size=9.5, pad=0.14, anchor=MSO_ANCHOR.MIDDLE, radius=0.08)
notes(s, "사용자가 고친 예전 판(진행자 → Gemini → 결과 아이콘 흐름)을 그대로 쓰고, 내용은 10/11 배포본(revision 00028)에서 실제로 넣은 문장과 화면 결과로 바꿨다. "
         "문장은 텍스트로 넣었다(음성이 아님, 화면에 '(typed)' 표시). 시간은 화면에 표시된 Shown in 값.")

# ---------------- 3 Opportunities ----------------
s = S[2]
strip_body(s, "Opportunities")
headline(s, "Checks against the regulation text, and the host fixes it on air")
label(s, 1.55, "Why now")


def stat_card(x, w, title, items, y=1.8, h=1.02):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill="FFFFFF", border=LINE, radius=0.06)
    tb(s, x + 0.14, y + 0.07, w - 0.28, 0.2, [[(title, {"bold": True})]], size=8.5, color=SUB)
    inner = w - 0.28 - (len(items) - 1) * 0.16
    ix = x + 0.14
    for num, text in items:
        iw = inner / len(items)
        tb(s, ix, y + 0.25, iw, 0.36, [[(num, {"bold": True})]], size=19, color=INK)
        tb(s, ix, y + 0.6, iw, h - 0.62, [text], size=8, color=SUB)
        ix += iw + 0.16


(c1, w3), (c2, _), (c3, _) = cols(3)
stat_card(c1, w3, "Market", [("25%", [("video commerce share of SEA e-commerce (2025, forecast)", {}), ("1", SUP)]),
                             ("~600M", [("live-commerce users in China (2024)", {}), ("2", SUP)])])
stat_card(c2, w3, "Risky claims on air", [("30 / 120", [("Korean live shows with suspected unfair ad claims (2020)", {}), ("3", SUP)])])
stat_card(c3, w3, "Rules now ask for it", [("2026", [("China: misstatements said on air must be corrected on the spot (Art. 31)", {}), ("4", SUP)])])
label(s, 2.95, "How different is it from any of the other existing ideas?")
table(s, 3.17, [(0.4, 2.3), (2.7, 2.9), (5.6, 4.0)], ["Common approach", "What it leaves out", "Ours"], [
    ["Keyword lists of banned words", "Words, not meaning: negation, rephrasing, quotes", "Gemini reads the whole line against the regulation PDF itself"],
    ["Compliance review after the show", "The claim has already gone out", "A red box in about 1-4 s, with a line to say instead"],
    ["Approving the script before air", "Hosts improvise; prices change", "Checks unscripted speech and today's prices, live"]], 0.25, size=8.5, head_size=8.5)
(b1, w2), (b2, _) = cols(2)
box(s, b1, 4.24, w2, 0.95, [[("How will it be able to solve the problem?", {"bold": True, "color": BLUE, "size": 9.5})],
    [("For brands and platforms that run their own live shows. The host reads a red or orange box from a distance, says the line "
      "shown, and the box clears. A solo host needs no operator; an operator can watch the same show.", {})]],
    fill="FFFFFF", border=LINE, size=8.5, pad=0.12, space=1, radius=0.06)
box(s, b2, 4.24, w2, 0.95, [[("USP of the proposed solution", {"bold": True, "color": BLUE, "size": 9.5})],
    [("Models judge meaning; code never overrules them. ", {"bold": True}),
     ("Any product enters through its own documents, with no product-specific prompts, and the correction is checked "
      "in the host's own words.", {})]],
    fill=BLUEBG, border=BLUE, size=8.5, pad=0.12, space=1, radius=0.06)
L1 = "https://blog.google/company-news/inside-google/around-the-globe/google-asia/sea-economy-2025/"
L2 = "https://www.cac.gov.cn/2025-07/28/c_1755417344368120.htm"
L3 = "https://www.kca.go.kr/smartconsumer/board/download.do?menukey=7301&fno=10030206&bid=00000146&did=1003136730"
L4 = "https://www.cac.gov.cn/2026-01/07/c_1769516655383334.htm"
tb(s, 0.4, 5.2, 9.2, 0.3, [
    [("1 ", SUP), ("Google, Temasek, Bain, e-Conomy SEA 2025 (forecast): ", FN), ("blog.google/…/sea-economy-2025", {"link": L1}),
     ("   2 ", SUP), ("CFIS & CIECC, China New E-commerce Development Report 2025: ", FN), ("cac.gov.cn/2025-07/28", {"link": L2})],
    [("3 ", SUP), ("Korea Consumer Agency, Live Commerce Advertising Survey (2020): ", FN), ("kca.go.kr (PDF)", {"link": L3}),
     ("   4 ", SUP), ("SAMR & CAC, Livestream E-commerce Supervision Measures, Art. 31 (2026): ", FN), ("cac.gov.cn/2026-01/07", {"link": L4})]],
   size=7, color=BLUE)
notes(s, "예전 제출본의 시장·위험 숫자와 출처(각주 1~4)를 그대로 썼다. 중국 규정 제31조(방송 중 말실수는 현장에서 바로잡는다)가 '말로 정정하면 경고가 사라지는' 흐름의 근거다. "
         "비교표에는 실제 회사 이름을 넣지 않았다(근거 없는 경쟁 비교를 피함, GPT 의견).")

# ---------------- 4 Features ----------------
s = S[3]
headline(s, "Before air people approve, on air the host only speaks, after air a recap")
feat = [
    ("Before air", "blue", ["Upload a product PDF; no fixed template to fill in", "Gemini reads it into a product record: brand, category, prices, prompter points, speech terms",
                            "Upload the regulation PDFs", "A person approves; nothing else is typed"]),
    ("On air: what the host says", "red", ["Red box: a claim a regulation clause forbids, the clause, a line to say instead",
                                          "Orange box: a price or fact that differs from the prompter",
                                          "Boxes clear when the host corrects it in any words",
                                          "Prompter ticks a point only when its full meaning is said; \"next\" + name or category moves to the next product"]),
    ("On air: what viewers ask", "green", ["Spam, abuse and prompt attacks screened (Model Armor)",
                                           "Answered automatically when the product record answers it, marked \"AI answer\"",
                                           "Health, claims and test results wait for the host or the operator, with a suggested reply"]),
]
for (x, w), (title, tone, items) in zip(cols(3), feat):
    fg, bg = TONE[tone]
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.68, w, 2.95, fill="FFFFFF", border=fg, radius=0.05, lw=1.25)
    shape(s, MSO_SHAPE.RECTANGLE, x + 0.01, 1.69, w - 0.02, 0.4, fill=bg)
    tb(s, x + 0.16, 1.69, w - 0.3, 0.4, [[(title, {"bold": True})]], size=11.5, color=fg, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x + 0.16, 2.22, w - 0.32, 2.35, [[("•  ", {"color": fg, "bold": True}), (t, {})] for t in items], size=10.5, space=7)
box(s, 0.4, 4.82, 9.2, 0.5, [[("After air: ", {"bold": True}), ("an AI recap from the show's numbers (alerts and how long each took to fix, "
                                                                 "unanswered questions, missed points), exported as PDF or JSON.", {})]],
    fill=BUBBLE, size=9.5, pad=0.14, anchor=MSO_ANCHOR.MIDDLE, radius=0.1)
notes(s, "세 칸: 방송 전(다른 팀: 제품·규정 업로드와 승인), 방송 중 진행자 말(우리 팀: STT 판정), 방송 중 시청자 질문(다른 팀: 댓글). 아래 줄: 방송 후 요약.")

# ---------------- 5 Process flow ----------------
s = S[4]
headline(s, "One approval before air sets the evidence every judge uses")
GX, GW = 1.55, 8.05
GG = 0.14
CW = (GW - 4 * GG) / 5


def lane(y, name, cells, h=0.62, tone="blue"):
    tb(s, 0.4, y, 1.1, h, [[(name, {"bold": True})]], size=9.5, color=TONE[tone][0], anchor=MSO_ANCHOR.MIDDLE)
    for i, (txt, sub, kind) in enumerate(cells):
        x = GX + i * (CW + GG)
        fill = {"ai": BLUEBG, "person": GRNBG, "place": "FFFFFF", "red": REDBG}[kind]
        color = {"ai": BLUE, "person": GRN, "place": SUB, "red": RED}[kind]
        box(s, x, y, CW, h, [[(txt, {"bold": True, "color": INK})], [(sub, {"color": color})]], fill=fill,
            border=LINE if kind == "place" else None, size=8.5, pad=0.08, anchor=MSO_ANCHOR.MIDDLE, radius=0.1)
        if i:
            arrow(s, x - GG + 0.01, y + h / 2, GG - 0.02)


lane(1.62, "Before air", [("Product PDF", "uploaded", "place"), ("Read into a record", "Gemini 3.8 Flash", "ai"),
                          ("Approve", "a person", "person"), ("Regulation PDFs", "uploaded as files", "place"),
                          ("Release", "product + regulations, in run order", "place")])
lane(2.42, "Host speech", [("Host mic", "browser, WebSocket", "place"), ("Speech to text", "Chirp 3, speech terms as hints", "ai"),
                           ("Final line", "a split sentence is merged", "place"), ("4 judges at once", "Gemini 3.5 Flash-Lite, 3.8 Flash", "ai"),
                           ("Host + operator", "red / orange box, ticks, switch", "red")])
lane(3.22, "Correction", [("Next line", "in the host's own words", "place"), ("Same judges", "is this alert corrected?", "ai"),
                          ("If corrected", "box clears: \"Cleared in N s\"", "person"), ("If not", "box stays; operator can cue or post a notice", "red"),
                          ("Alert log", "every verdict kept", "place")], tone="red")
lane(4.02, "Viewer chat", [("Comment", "viewer page", "place"), ("Screen", "code filter + Model Armor", "ai"),
                           ("Triage", "Gemini 3.8 Flash", "ai"), ("Answer from records", "Vertex AI Search, then checked", "ai"),
                           ("Post or wait", "AI answer, or host / operator", "person")], tone="green")
tb(s, 0.4, 4.8, 9.2, 0.55, [[("Legend  ", {"bold": True, "color": SUB}), ("blue = Gemini", {"color": BLUE}), ("   green = a person or the result they see", {"color": GRN}),
                             ("   red = what the host must fix", {"color": RED})],
                            [("The judges read only what was approved: the regulation files for red, what the prompter shows for orange. "
                              "Nothing is typed in by hand during the show.", {"color": SUB})]], size=8.5, space=3)
notes(s, "네 줄이 같은 다섯 칸 격자를 쓴다. 파랑 = Gemini, 초록 = 사람 또는 결과, 빨강 = 진행자가 고칠 것. 판정 AI는 승인된 것만 읽는다(빨강: 규정 파일, 주황: 프롬프터 내용).")

# ---------------- 6 Wireframes ----------------
s = S[5]
headline(s, "Alerts, the line to say and the talking points on one screen")
WX, WY, WW, WH = 0.4, 1.62, 6.0, 3.72
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, WX, WY, WW, WH, fill="FFFFFF", border=LINE, radius=0.03)
tb(s, WX + 0.36, WY + 0.08, 1.4, 0.24, [[("● LIVE 00:12:31", {"bold": True, "color": RED})]], size=9)
tb(s, WX + 1.75, WY + 0.06, 3.0, 0.3, [[("NOW ", {"color": SUB, "size": 7}), ("Hydra Barrier Cream", {"bold": True})],
                                      [("NEXT ", {"color": SUB, "size": 7}), ("Deep Hydration Ampoule", {"color": SUB})]], size=8.5)
hline(s, WX, WY + 0.44, WX + WW)
PX, PW = WX + 0.12, 3.9
box(s, PX, WY + 0.55, PW, 0.72, [[("Regulation", {"bold": True, "color": RED}), ("   Heard: \"treats eczema\"", {"color": SUB, "size": 7.5})],
                                  [("Take it back, then say:", {"color": RED, "size": 7.5})],
                                  [("\"Sorry, I take that back. Helps relieve the feeling of dryness.\"", {"bold": True})],
                                  [("ASEAN guideline App. III 2.d · Shown in 1.3 s", {"color": SUB, "size": 7})]],
    fill=REDBG, border=RED, size=8.5, pad=0.1, radius=0.06)
box(s, PX, WY + 1.35, PW, 0.46, [[("Product info", {"bold": True, "color": YELT}), ("   Heard: S$32", {"color": SUB, "size": 7.5})],
                                  [("The live price is S$36.00.", {"bold": True})]],
    fill=YELBG, border=YELT, size=8.5, pad=0.1, radius=0.06)
tb(s, PX, WY + 1.9, 1.1, 0.5, [[("LIVE PRICE", {"color": SUB, "size": 7})], [("S$36.00", {"bold": True, "size": 14})]], size=8)
tb(s, PX + 1.2, WY + 1.9, PW - 1.2, 0.8, [[("✓ ", {"color": GRN, "bold": True}), ("Hydrated for up to 24 hours in our in-house test", {})],
                                           [("★ ", {"color": GRN, "bold": True}), ("Suitable for sensitive skin, patch tested", {"bold": True})],
                                           [("•  ", {"color": SUB}), ("Helps support the skin's moisture barrier", {})]], size=7.5, space=2)
hline(s, PX, WY + 2.78, PX + PW)
tb(s, PX, WY + 2.84, PW, 0.8, [[("QUESTIONS", {"color": SUB, "size": 7})],
                                [("01 Is it suitable for sensitive skin that gets redness easily?", {"bold": True})],
                                [("It is suitable for sensitive skin. If redness appears, stop and consult a dermatologist.", {})]], size=7.5, space=2)
CX = WX + 4.15
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, CX, WY + 0.55, WW - 4.27, 3.05, fill=BUBBLE, radius=0.04)
tb(s, CX + 0.1, WY + 0.62, 1.6, 2.9, [[("LIVE CHAT", {"color": SUB, "size": 7})], [("Mei ", {"bold": True}), ("I get redness easily...", {})],
                                       [("Arjun ", {"bold": True}), ("how much is it today?", {})],
                                       [("AI answer ", {"bold": True, "color": BLUE}), ("The live price is S$36.00...", {"color": SUB})],
                                       [("Siti ", {"bold": True}), ("does it have fragrance?", {})]], size=7, space=4)
callouts = [(WY + 0.15, "Show clock and the product on air, with the next one"),
            (WY + 0.85, "Red: the regulation clause and one full sentence to read"),
            (WY + 1.52, "Orange: the right value from the prompter"),
            (WY + 2.2, "Prompter: ticks fill in as points are said"),
            (WY + 3.1, "Viewer questions that need a person, with a reply")]
for k, (cy, text) in enumerate(callouts):
    num_circle(s, 6.62, cy, k + 1)
    tb(s, 7.0, cy - 0.03, 2.6, 0.5, [text], size=10, anchor=MSO_ANCHOR.TOP)
for k, (x, y) in enumerate([(WX + 0.06, WY + 0.04), (PX + PW - 0.2, WY + 0.5), (PX + PW - 0.2, WY + 1.3),
                            (PX + PW - 0.2, WY + 1.9), (PX + PW - 0.2, WY + 2.84)]):
    num_circle(s, x, y, k + 1, d=0.22)
tb(s, 6.62, WY + 3.5, 3.0, 0.3, [[("Layout schematic. The host never clicks; real screens on slide 10.", {})]], size=8.5, color=SUB)
notes(s, "진행자 화면 구조도(실제 캡처는 10장). 번호 1~5: 방송 시계와 지금/다음 상품, 빨강(규정 조항과 읽을 한 문장), 주황(맞는 값), 프롬프터 체크, 사람이 필요한 시청자 질문과 추천 답변.")

# ---------------- 7 Architecture ----------------
s = S[6]
headline(s, "One Cloud Run service runs each show and calls managed Google Cloud AI")
label(s, 1.55, "Inputs and screens", x=0.4, w=2.2, color=SUB)
label(s, 1.55, "Cloud Run service (FastAPI)", x=2.95, w=3.6, color=BLUE)
label(s, 1.55, "Google Cloud managed services", x=6.85, w=2.75, color=SUB)
for k, (t1, t2) in enumerate([("Host mic", "audio over WebSocket"), ("Viewer chat", "comments"), ("Product + regulation PDFs", "uploaded before air"),
                               ("Screens", "host, operator, viewer, summary (SSE)")]):
    y = 1.85 + k * 0.82
    box(s, 0.4, y, 2.2, 0.62, [[(t1, {"bold": True})], [(t2, {"color": SUB, "size": 8.5})]], fill="FFFFFF", border=LINE, size=10,
        pad=0.12, anchor=MSO_ANCHOR.MIDDLE, radius=0.08)
    if k < 3:
        arrow(s, 2.65, y + 0.31, 0.25)
    else:
        shape(s, MSO_SHAPE.LEFT_ARROW, 2.65, y + 0.25, 0.25, 0.12, fill="9AA0A6")
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.95, 1.85, 3.6, 3.26, fill="FFFFFF", border=BLUE, radius=0.04, lw=1.5)
tb(s, 3.1, 1.93, 3.3, 0.4, [[("STT stream per show · a split sentence is merged · timers", {})]], size=8.5, color=SUB)
judges = [("Compliance", "Flash-Lite · regulation PDFs as files", "red"), ("Fact checker", "3.8 Flash · what the prompter shows", "yellow"),
          ("Context", "Flash-Lite · switch, question hold", "blue"), ("Prompter", "3.8 Flash · full-meaning ticks", "green")]
for k, (t1, t2, tone) in enumerate(judges):
    x = 3.1 + (k % 2) * 1.67
    y = 2.3 + (k // 2) * 0.72
    fg, bg = TONE[tone]
    box(s, x, y, 1.6, 0.62, [[(t1, {"bold": True, "color": fg})], [(t2, {"color": INK, "size": 7.5})]], fill=bg, size=9.5, pad=0.08,
        anchor=MSO_ANCHOR.MIDDLE, radius=0.08)
tb(s, 3.1, 3.76, 3.3, 0.24, [[("4 judges in parallel on every final line", {"bold": True})]], size=8.5, color=BLUE)
tb(s, 3.1, 4.02, 3.3, 1.0, [[("•  ", {"color": BLUE}), ("Comment path: screen, triage, answer, check", {})],
                            [("•  ", {"color": BLUE}), ("Code maps ids, runs timers, draws screens; it never drops a verdict", {})],
                            [("•  ", {"color": BLUE}), ("IAP login · one instance · session affinity", {})]], size=8.5, space=3)
shape(s, MSO_SHAPE.LEFT_RIGHT_ARROW, 6.6, 3.36, 0.22, 0.16, fill="9AA0A6")
svc = [("Speech-to-Text V2", "chirp_3, region us"), ("Vertex AI Gemini", "3.5 Flash-Lite · 3.8 Flash (global)"),
       ("Vertex AI Search", "product sheets for viewer answers"), ("Model Armor", "screens viewer comments"),
       ("Firestore", "show state, alerts, every AI call"), ("Cloud Storage", "products, regulations, releases (write-once)")]
for k, (t1, t2) in enumerate(svc):
    y = 1.85 + k * 0.555
    box(s, 6.85, y, 2.75, 0.48, [[(t1, {"bold": True}), ("  " + t2, {"color": SUB, "size": 7.5})]], fill="FFFFFF", border=LINE,
        size=9, pad=0.1, anchor=MSO_ANCHOR.MIDDLE, radius=0.1)
tb(s, 0.4, 5.2, 9.2, 0.25, [[("Each judge has one source of truth: regulation files for red, the prompter for orange. "
                              "Show state and every AI call are stored in Firestore.", {})]], size=8.5, color=SUB)
notes(s, "예전 3단 구조(입력·출력 | Cloud Run | 관리형 서비스)를 그대로 쓰고 내용을 바꿨다. 가운데 Cloud Run 안에 판정 AI 4개(규정·사실·문맥·프롬프터)를 색으로 구분했다.")

# ---------------- 8 Technologies ----------------
s = S[7]
headline(s, "Each Google Cloud service has one job in the show")
table(s, 1.6, [(0.4, 2.25), (2.65, 2.85), (5.5, 4.1)], ["Service", "Model or setting", "Its job"], [
    [[("Speech-to-Text V2", {"bold": True})], "chirp_3, region us, STANDARD endpointing", "Streams the host's speech; speech terms from the product record as hints"],
    [[("Vertex AI Gemini", {"bold": True})], "3.5 Flash-Lite", "Compliance judge and context judge (one call each per line)"],
    [[("Vertex AI Gemini", {"bold": True})], "3.8 Flash", "Fact checker and prompter judge; reading product PDFs; comment triage, answers, recap"],
    [[("Vertex AI Search", {"bold": True})], "product sheets, per market", "Finds the product-sheet passage behind a viewer answer"],
    [[("Model Armor", {"bold": True})], "before any model", "Screens viewer comments for spam, abuse and prompt attacks"],
    [[("Cloud Run + IAP", {"bold": True})], "us-central1, one instance", "FastAPI backend, React screens, server-sent events"],
    [[("Firestore", {"bold": True})], "Native mode", "Sessions, transcripts, alerts, segments and every AI call with its tokens"],
    [[("Cloud Storage", {"bold": True})], "write-once objects", "Approved products, regulation files and releases"],
    [[("App stack", {"bold": True})], "Python 3.12, FastAPI, React 18, TypeScript, Vite", "Agent Development Kit for the AI PD agent; pytest, Vitest, Playwright"]],
    0.38, size=9)
notes(s, "로고 나열 대신 '서비스 | 설정 | 맡은 일' 표(GPT·Gemini 공통 의견).")

# ---------------- 9 Cost ----------------
s = S[8]
headline(s, "Judges + speech: about US$4.2 an hour at 600 lines an hour")
table(s, 1.6, [(0.4, 2.0), (2.4, 4.3), (6.7, 1.4), (8.1, 1.5)], ["Item", "Basis (measured)", "Per line", "Per hour (600 lines)"], [
    ["Compliance judge", "Flash-Lite, ~16.8k tokens (regulation PDFs), cached", "$0.0011", "$0.63"],
    ["Context judge", "Flash-Lite, ~4.0k tokens in", "$0.0016", "$0.94"],
    ["Prompter judge", "3.8 Flash, ~1.2k tokens in, low thinking", "$0.0012", "$0.74"],
    ["Fact checker", "3.8 Flash, ~5.3k tokens, part cached (est.)", "$0.0016", "$0.95"],
    ["Speech-to-Text", "chirp_3, $0.016 per audio minute", "", "$0.96"],
    [[("Total", {"bold": True})], "judges per line; per hour adds speech", [("$0.0054", {"bold": True})], [("≈ $4.2", {"bold": True})]]],
    0.36, size=9.5)
(b1, w2), (b2, _) = cols(2)
box(s, b1, 4.22, w2, 1.0, [[("How it was measured", {"bold": True, "color": BLUE})],
                            ["Token counts of every judge call in 2 shows on the deployed app (10/11, 14 host lines) from Firestore, "
                             "times Google's list prices. The fact checker moved to 3.8 Flash on 10/15; its row is estimated from the same tokens. Per hour assumes 10 final lines a minute (600 lines, an assumption, not measured)."]],
    fill="FFFFFF", border=LINE, size=8.5, pad=0.12, space=2, radius=0.06)
box(s, b2, 4.22, w2, 1.0, [[("Not included", {"bold": True, "color": YELT})],
                            ["Repeat requests sent when a judge is slow (not logged; at most doubles that call). Viewer comments: ~$0.006 a comment "
                             "for triage. The recap: $0.007 a show. Cloud Run and Firestore. 3.8 Flash list price doubles on 2027-01-01."]],
    fill=YELBG, size=8.5, pad=0.12, space=2, radius=0.06)
tb(s, 0.4, 5.3, 9.2, 0.18, [[("List prices on 10/11: Gemini API pricing page (ai.google.dev) and Speech-to-Text pricing (cloud.google.com).", {})]], size=7, color=SUB)
notes(s, "비용: Firestore agent_calls의 토큰(2개 방송, 진행자 문장 14개, 10/11 배포본) × 공식 단가. 3.5 Flash-Lite 입력 $0.30, 캐시 $0.03, 출력 $2.50, "
         "3.8 Flash 입력 $0.75, 출력 $3.75 (12/31까지, 2027년 2배). STT Standard(chirp 포함) $0.016/분. 1시간 600문장은 가정. "
         "헤지(느린 호출에 두 번째 요청)는 기록되지 않아 빠졌다. 사실 확인은 10/15에 3.8 Flash로 바뀌어 같은 토큰으로 추정했다(D-466: 시간당 약 +$0.5).")

# ---------------- 10 Snapshots ----------------
s = S[9]
headline(s, "One show on the deployed app, from alert to correction")
shots = [("alerts", "1  Red and orange boxes on the host screen"), ("cleared", "2  Both cleared in the host's words; points ticked"),
         ("operator", "3  Operator: the clause, the AI summary, actions"), ("questions", "4  Viewer questions that need the host, with replies")]
for k, (name, cap) in enumerate(shots):
    w = 3.9
    x = (0.75, 5.35)[k % 2]
    path = IMG / f"{name}.png"
    pw, ph = Image.open(path).size
    h = w * ph / pw
    y = 1.58 + (k // 2) * (h + 0.38)
    picture(s, name, x, y, w=w, h=h)
    tb(s, x, y + h + 0.05, w, 0.28, [cap], size=9.5, color=INK)
tb(s, 0.4, 5.3, 9.2, 0.2, [[("Cloud Run revision 00028, 10/11. Lines typed into the live judging path (marked \"typed\"); "
                             "speech goes through Chirp 3 to the same judges (measured on slide 11).", {})]], size=8, color=SUB)
notes(s, "10/11 배포본에서 만든 한 방송의 캡처 4장(검토 의견: 6장은 글씨가 안 보여 4장으로 키움). 진행자 문장은 텍스트로 넣었다(음성 아님). 원본 1920x1080 캡처는 스크래치패드에 있다.")

# ---------------- 11 Performance ----------------
s = S[10]
headline(s, "Detection 96.3% and no false alerts on a fresh set; fact checks are slow")
kpis = [("Detection", "96.3%", "52 of 54 · other settings 100%", "≥ 95%", "green"),
        ("False alerts on safe lines", "0%", "0 of 108 · other settings 0-0.9%", "< 1%", "green"),
        ("Compliance judge, one call p95", "2.9 s", "input to output, no retries", "≤ 5 s", "green"),
        ("Fact checker, one call p95", "31.5 s", "slow night · 5.3-12.9 s on 10/14", "≤ 5 s · not met", "red")]
for (x, w), (t, num, sub, line, tone) in zip(cols(4), kpis):
    fg, bg = TONE[tone]
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.6, w, 1.2, fill="FFFFFF", border=LINE, radius=0.06)
    tb(s, x + 0.14, 1.67, w - 0.28, 0.2, [[(t, {"bold": True})]], size=8.5, color=SUB)
    tb(s, x + 0.14, 1.88, w - 0.28, 0.42, [[(num, {"bold": True})]], size=22)
    tb(s, x + 0.14, 2.3, w - 0.28, 0.25, [sub], size=8, color=SUB)
    chip(s, x + 0.14, 2.52, "Pass line " + line, tone, size=7.5)
tb(s, 0.4, 2.88, 9.2, 0.38, [[("Sealed set 10: 4 products the app never saw, invented documents; 162 natural host lines in English and Korean spoken by a TTS voice through the real "
                               "Chirp 3 path (12 lines cut in two on purpose and 6 out-of-scope lines reported apart). Labels by GPT and a blind Claude critic. Red alert on a non-regulation line: 0. Per product 92.3-100%.", {})]],
   size=8, color=SUB)
(b1, w2), (b2, _) = cols(2)
label(s, 3.32, "Why models judge, not code", x=b1, w=w2)
table(s, 3.56, [(b1, 2.1), (b1 + 2.1, 1.25), (b1 + 3.35, 1.15)], ["Same STT lines (set 5)", "False alerts", "Wrong strongest"], [
    ["Code rules on top of the model", [("19.4%", {"color": RED, "bold": True})], "5"],
    ["Model verdicts only", [("0.9-1.9%", {"color": GRN, "bold": True})], "1"]], 0.3, size=9, head_size=8)
label(s, 4.56, "Detection on each new unseen set", x=b1, w=w2)
tb(s, b1, 4.8, w2, 0.3, [[("84% → 91.0% → 89.6-90.3% → 94.0% → ", {}), ("96.3%", {"bold": True})]], size=11)
tb(s, b1, 5.08, w2, 0.2, [["sets 3, 4, 6, 7, 10 (from set 6 through TTS and STT; earlier sets ran under older conditions)"]], size=7.5, color=SUB)
label(s, 3.32, "Other measured parts", x=b2, w=w2)
table(s, 3.56, [(b2, 2.75), (b2 + 2.75, 1.75)], ["Test", "Result"], [
    ["Prompter judge, partial vs full points (10/11 probe)", "wrong ticks 0 / 40, missed 0 / 63"],
    ["Product switch on a category word only", "0 / 5 → 4 / 4 with category"],
    ["Viewer comments: question / reaction / hide (233)", "99.6%"],
    ["Comments: auto vs needs a person, as labelled", "90.6% · injection replies 0"]], 0.32, size=8.5, head_size=8)
tb(s, 0.4, 5.32, 9.2, 0.2, [[("Measured on the current judge setting with a TTS voice; no real human host yet. A slow fact check now shows late instead of being dropped.", {})]],
   size=8.5, color=RED)
notes(s, "README의 측정값 그대로(docs/eval/sealed-set-10-results.md). 세트 10, 지금 설정(규정 Flash-Lite LOW, 사실 3.8 Flash LOW), Gemini가 느리던 10/15 밤: 탐지 52/54, 놓친 2줄은 사실 확인 답이 끝내 오지 않은 줄. 다른 규정 설정 3개(10/14)는 54/54. "
         "코드 규칙 vs 모델 판정(세트 5, 진단용). 댓글 수치는 다른 팀의 docs/comments.md(1회, 233개 댓글).")

# ---------------- 12 Additional details / future ----------------
s = S[11]
headline(s, "Flags people act on; next, faster fact checks and a larger test")


def item_box(x, w, title, items, kind, tone):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.62, w, 3.7, fill="FFFFFF", border=LINE, radius=0.05)
    tb(s, x + 0.18, 1.72, w - 0.36, 0.3, [[(title, {"bold": True})]], size=11, color=TONE[tone][0])
    for k, t in enumerate(items):
        icon(s, kind, x + 0.18, 2.12 + k * 0.62, 0.24)
        tb(s, x + 0.54, 2.1 + k * 0.62, w - 0.72, 0.58, [t], size=9.5)


(b1, w2), (b2, _) = cols(2)
item_box(b1, w2, "Responsible AI (in the prototype)", [
    "A red box marks a claim for the host to fix. It is not a legal ruling.",
    "The line to say instead is built only from the product's approved claims.",
    "Health, claim and test-result questions wait for a person; automatic replies say \"AI answer\".",
    "Judge calls are stored with input, output and tokens; code never drops a verdict.",
    "No product-specific prompts: a new product enters only through its documents."], "check", "green")
item_box(b2, w2, "Next (planned, not in the prototype)", [
    "Faster fact checks: a quicker model when Gemini 3.8 Flash is slow.",
    "Test with a real human host and two hosts speaking over each other.",
    "A larger sealed set (set 10 had 4 products and 54 violations).",
    "Log repeat requests so the cost per show is complete.",
    "Copy AI-call logs to BigQuery for analysis across shows."], "plus", "blue")
notes(s, "왼쪽(초록 체크 = 지금 있음): 법적 판단이 아님, 대체 문구는 승인된 문구에서만, 사람 확인 질문, 모든 AI 호출 기록, 제품별 프롬프트 없음. "
         "오른쪽(파랑 더하기 = 계획): 느린 날 사실 확인 속도, 사람 진행자 시험, 더 큰 봉인 세트, 비용 기록 보완, BigQuery.")

# ---------------- 13 Links ----------------
s = S[12]
strip_body(s, "Provide links to your")
headline(s, "Try it: start a show, say a claim, then take it back")
links = [("GitHub Public Repository", "github.com/jwchoiKR/ai_builder_cup_uh", "https://github.com/jwchoiKR/ai_builder_cup_uh"),
         ("Demo Video Link (3 Minutes)", None, None), ("Final Product Link", None, None)]
for (x, w), (name, text, url) in zip(cols(3), links):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.65, w, 0.95, fill="FFFFFF", border=LINE, radius=0.07)
    tb(s, x + 0.18, 1.76, w - 0.36, 0.3, [[(name, {"bold": True})]], size=11)
    if url:
        tb(s, x + 0.18, 2.14, w - 0.36, 0.3, [[(text, {"link": url})]], size=9.5)
    else:
        chip(s, x + 0.18, 2.14, "to be added", "yellow", size=9)
label(s, 2.85, "Three steps for an evaluator")
steps = ["Open the hub and start a show from an approved product",
         "Say \"This cream treats eczema\" and watch the red box",
         "Read the line it shows and watch the box clear"]
for k, ((x, w), text) in enumerate(zip(cols(3), steps)):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.15, w, 0.85, fill="FFFFFF", border=BLUE, radius=0.07, lw=1.5)
    num_circle(s, x + 0.16, 3.43, k + 1)
    tb(s, x + 0.56, 3.15, w - 0.7, 0.85, [[(text, {"bold": True})]], size=10, anchor=MSO_ANCHOR.MIDDLE)
    if k:
        arrow(s, x - 0.18, 3.57, 0.16)
tb(s, 0.4, 4.2, 9.2, 0.5, [["Every alert comes from the live models on the same path as a real show; nothing is pre-stored. "
                            "Speak into the browser mic on the host screen."]], size=9.5, color=SUB)
notes(s, "링크 세 칸(양식 문구 그대로). 데모 영상과 제품 링크는 아직 없다. 제품 링크는 지금 IAP 로그인 뒤에 있어서, 심사위원이 열 수 있게 할지 정해야 한다. "
         "체험 3단계: 쇼 시작 → 위반 문장 말하기 → 보여 준 문장 읽고 경고가 사라지는 것 보기.")

prs.save(OUT)
print(OUT)
